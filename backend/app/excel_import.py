"""历史周会 Excel 导入（v3 Phase-17，领导 2026-09-30："把上传功能做成真的，导入历史数据、替换演示数据"）。

输入是委员会每周的 Agenda / MM 文件（SharePoint 文档库里一周一个文件夹），2025-09 至今的 64 份文件
列结构一致：第 1 行表头，'#' / Region / Type / Project / Commodity / Parent PF / SWAT ID / Part Description / PN /
Recommended Supplier / Presenter / Decision Level / Peak Year Spend / LT Spend / Comment / Sourcing Decision / Notes /
CQA PC PRICE / Rec. Supplier Landed / Average Volume / LT volume / CQA Tooling / Rec. Supplier Tooling。
一个案例占若干行（案例级单元格纵向合并），零件按行；金额多为文本（"439K Euro"、"2,5Mio Euro⏎2.7 Mio USD"）。

解析规则（都有单元测试）：
- 表头行 = 同时含 "PN" 与 "Region" 的那一行；列按去掉空格标点后的前缀匹配，多一列少一列都不报错。
- 合并单元格先展开：每个格子记住自己是"合并填充"来的还是"自己的值"；案例边界 = '#' 是自己的值的那一行。
- 金额：取第一个带欧元单位的数（K / Mio / M 倍数），只有 USD / RMB 时按 Dashboard 汇率折算，都没有则记 0 并给警告；
  逗号 / 点号：件价里逗号一律当小数点；总额里 "60,000" 这种 3 位一组当千分位、"1,03 Mio" 当小数。
- 决议文案按关键词归类（rejected / not approved → REJECTED；condition / with action(s) → CONDITIONAL APPROVAL；
  approved → APPROVED；其余 → PENDING），原文保留在 decisionText。
- 会议日期与 KW 从文件名取（"KW39 23.09.2026"、"KW37, 10.09.2025"、"20260128 … KW05"、"WK15. Apr 08 2026"），
  调用方可覆盖；同一年同一周再次导入时替换上次导入的案例。
"""

from __future__ import annotations

import io
import re
from datetime import date
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import logic

SOURCE_KIND = "excel"
MAX_FILE_BYTES = 20 * 1024 * 1024

# 表头别名：去掉空格与标点、小写后匹配（'#' 与 'pn' 精确匹配，其余按前缀）
COLUMN_KEYS: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("num", ("#",)),
    ("region", ("region",)),
    ("type", ("type",)),
    ("project", ("project",)),
    ("commodity", ("commodity",)),
    ("parentPF", ("parentpf",)),
    ("swat", ("swatid", "swat")),
    ("partDescription", ("partdescription",)),
    ("pn", ("pn",)),
    ("supplier", ("recommendedsupplier",)),
    ("presenter", ("presenter",)),
    ("decisionLevel", ("decisionlevel",)),
    ("peak", ("peakyearspend",)),
    ("lifetime", ("ltspend", "lifetimespend")),
    ("comment", ("comment",)),
    ("decision", ("sourcingdecision",)),
    ("notes", ("notes",)),
    ("pcPriceCQA", ("cqapcprice",)),
    ("supplierPriceLanded", ("recsupplierlanded",)),
    ("averageVolume", ("averagevolume",)),
    ("lifetimeVolume", ("ltvolume", "lifetimevolume")),
    ("toolingCQA", ("cqatooling",)),
    ("supplierToolingCost", ("recsuppliertooling",)),
)
EXACT_KEYS = {"#", "pn"}
FIELDS = tuple(f for f, _ in COLUMN_KEYS)
ROW_MONEY = ("pcPriceCQA", "supplierPriceLanded", "toolingCQA", "supplierToolingCost")
ROW_INT = ("averageVolume", "lifetimeVolume")
PART_FIELDS = {"pn", "partDescription", "swat", *ROW_MONEY, *ROW_INT}

AMOUNT_RE = re.compile(
    r"(?P<num>\d[\d.,]*)\s*(?P<mult>mio|mil|m|k|tsd)?\s*(?P<unit>euro|eur|€|usd|\$|rmb|cny|¥)",
    re.IGNORECASE,
)
NUMBER_RE = re.compile(r"-?\d[\d.,]*")
BARE_MULTIPLIER_RE = re.compile(r"(?P<num>\d[\d.,]*)(?P<mult>mio|mil|m|k|tsd)", re.IGNORECASE)
BUNDLE_RE = re.compile(r"B/\d{4}/\d{5}", re.IGNORECASE)
SWAT_RE = re.compile(r"\b\d{9,12}\b")
PN_TOKEN_RE = re.compile(r"^[A-Za-z0-9-]{5,16}$")
MULTIPLIER = {"k": 1_000, "m": 1_000_000, "mio": 1_000_000, "mil": 1_000_000, "tsd": 1_000}
UNIT_CURRENCY = {"euro": "EUR", "eur": "EUR", "€": "EUR", "usd": "USD", "$": "USD", "rmb": "CNY", "cny": "CNY", "¥": "CNY"}
EMPTY_TEXT = {"", "-", "—", "n/a", "tbd", "none", "null"}  # 注意 "NA" 是北美区域，不是空值
CASE_STATUS = {"APPROVED": "Approved", "CONDITIONAL APPROVAL": "Approved", "REJECTED": "Resubmission Required"}


class TemplateError(ValueError):
    """文件不是委员会模板、猜不出会议日期等：给管理员看的可读错误。"""


# ---------- 文件名 → 周 / 日期 ----------

_MONTHS = {m: i for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"), 1)}


def meeting_from_name(name: str) -> tuple[int | None, date | None]:
    """从文件名（或 SharePoint 周文件夹名）猜 (KW, 会议日期)；猜不到的项为 None。

    文件名里的日期有时是准备日（"KW39 24 26.09.2025"、"KW36 31.08.2026"）：KW 与日期不一致时以 KW 为准，取那周的周三。"""
    stem = re.sub(r"\.xlsx?$", "", name, flags=re.IGNORECASE)
    kw = None
    if (m := re.search(r"(?<![A-Za-z0-9])(?:KW|WK)\s*(\d{1,2})(?!\d)", stem, re.IGNORECASE)):  # "KW46_SBS" 下划线紧跟也算
        kw = int(m.group(1))
    day: date | None = None
    # 文件名里日期后面常紧跟下划线（"23.09.2026_SBS"），下划线算单词字符，所以不用 \b 而用前后不是数字来界定
    if (m := re.search(r"(?<!\d)(\d{1,2})\.(\d{1,2})\.(20\d{2})(?!\d)", stem)):  # 23.09.2026
        day = _safe_date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    elif (m := re.search(r"(?<!\d)(20\d{2})(\d{2})(\d{2})(?!\d)", stem)):  # 20260128
        day = _safe_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    elif (m := re.search(r"\b([A-Za-z]{3})[a-z]*\.?\s+(\d{1,2})\s+(20\d{2})\b", stem)):  # Apr 08 2026
        month = _MONTHS.get(m.group(1).lower())
        day = _safe_date(int(m.group(3)), month, int(m.group(2))) if month else None
    if day is None and kw is not None and (m := re.search(r"(?<!\d)(20\d{2})(?!\d)", stem)):
        day = wednesday_of_iso_week(int(m.group(1)), kw)
    if day is not None and kw is None:
        kw = day.isocalendar().week
    if day is not None and kw is not None:
        # 委员会周三开会；文件名里的日期可能是周一或周五的准备日 —— 以 KW 为准取那周的周三（跨年周按离得近的年份）
        candidates = {wednesday_of_iso_week(y, kw) for y in (day.isocalendar().year, day.year) if 1 <= kw <= 53}
        day = min(candidates, key=lambda d: abs((d - day).days)) if candidates else day
    return kw, day


def wednesday_of_iso_week(year: int, week: int) -> date:
    return date.fromisocalendar(year, week, 3)


def _safe_date(y: int, m: int | None, d: int) -> date | None:
    try:
        return date(y, m or 0, d)
    except ValueError:
        return None


def meeting_label(day: date) -> str:
    """与提交单一致的日期文案："Wed, Sep 30, 2026"。"""
    return f"{day:%a}, {day:%b} {day.day}, {day.year}"


# ---------- 单元格值 ----------

def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", "\n")
    text = re.sub(r"[ \t\xa0]+", " ", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def _line(value: Any) -> str:
    return _clean(value).replace("\n", " ")


def _is_blank(value: Any) -> bool:
    return _clean(value).lower() in EMPTY_TEXT


def _to_decimal(num: str, *, price: bool, has_multiplier: bool) -> Decimal | None:
    """'2,5' / '1.03' / '60,000' / '253,164.56' / '370.000' → Decimal。"""
    s = num.strip()
    if not s:
        return None
    if "," in s and "." in s:  # 两种分隔符都有：最后一个是小数点
        last = max(s.rfind(","), s.rfind("."))
        s = re.sub(r"[.,]", "", s[:last]) + "." + s[last + 1:]
    elif price:  # 件价从不写千分位：逗号就是小数点
        s = s.replace(",", ".")
        if s.count(".") > 1:
            s = s.replace(".", "")
    else:
        sep = "," if "," in s else "." if "." in s else ""
        if sep:
            parts = s.split(sep)
            thousands = (len(parts) > 2
                         or (len(parts) == 2 and len(parts[1]) == 3 and parts[0] != "0" and not has_multiplier))
            s = "".join(parts) if thousands else parts[0] + "." + "".join(parts[1:])
    try:
        return Decimal(s)
    except InvalidOperation:
        return None


def parse_amount(value: Any, fx_settings: dict | None, *, price: bool = False) -> tuple[Decimal | None, str | None]:
    """金额格子 → (欧元数, 警告)。数字按欧元；文本找第一个带欧元单位的数，只有外币时按 Dashboard 汇率折算。"""
    if value is None or isinstance(value, bool) or _is_blank(value):
        return None, None
    if isinstance(value, (int, float)):
        return Decimal(str(value)), None
    text = _clean(value)
    candidates: list[tuple[str, Decimal]] = []
    for m in AMOUNT_RE.finditer(text):
        n = _to_decimal(m.group("num"), price=price, has_multiplier=bool(m.group("mult")))
        if n is None:
            continue
        if m.group("mult"):
            n *= MULTIPLIER[m.group("mult").lower()]
        candidates.append((UNIT_CURRENCY[m.group("unit").lower()], n))
    for cur, n in candidates:
        if cur == "EUR":
            return n, None
    for cur, n in candidates:
        rate = _per_eur(fx_settings, cur)
        if rate is None:
            return None, f"{cur} amount but no {cur} rate on the Dashboard: {text[:40]!r}"
        return (n / rate).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP), f"{cur} converted at 1 EUR = {rate} {cur}"
    compact = text.replace(" ", "")
    if (m := BARE_MULTIPLIER_RE.fullmatch(compact)):  # "5k"、"2.27Mio"：没写币种，按欧元
        n = _to_decimal(m.group("num"), price=price, has_multiplier=True)
        if n is not None:
            return n * MULTIPLIER[m.group("mult").lower()], None
    if NUMBER_RE.fullmatch(compact):  # 没有单位、整格就是一个数：按欧元
        n = _to_decimal(compact.lstrip("-"), price=price, has_multiplier=False)
        if n is not None:
            return (-n if compact.startswith("-") else n), None
    return None, f"cannot read amount {text[:40]!r}"


def _per_eur(fx_settings: dict | None, currency: str) -> Decimal | None:
    raw = ((fx_settings or {}).get("perEur") or {}).get(currency)
    try:
        value = Decimal(str(raw)) if raw not in (None, "") else None
    except InvalidOperation:
        return None
    return value if value and value > 0 else None


def parse_int(value: Any) -> int | None:
    if value is None or isinstance(value, bool) or _is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return round(value)
    m = NUMBER_RE.search(_clean(value).replace(" ", ""))
    if not m:
        return None
    n = _to_decimal(m.group(0).lstrip("-"), price=False, has_multiplier=False)
    return int(n) if n is not None else None


def _money_text(amount: Decimal | None) -> str | None:
    return logic.money_text(amount.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)) if amount is not None else None


def _whole_eur(amount: Decimal) -> int:
    return int(amount.quantize(Decimal(1), rounding=ROUND_HALF_UP))


# ---------- 文案归一 ----------

def classify_decision(text: str) -> str:
    t = _clean(text).lower()
    if not t or t in EMPTY_TEXT:
        return "PENDING"
    if re.search(r"reject|not approved|declined", t):
        return "REJECTED"
    if re.search(r"condition|with action", t):
        return "CONDITIONAL APPROVAL"
    if "approved" in t:
        return "APPROVED"
    return "PENDING"


def normalize_type(text: str) -> str:
    t = _line(text).lower().replace("-", " ")
    if not t:
        return ""
    if "gcs" in t:
        return "GCS"
    if re.search(r"carry|\bc/o\b|\bco\b", t):
        return "C/O"
    if "new" in t or "family" in t:
        return "New"
    return _line(text)


def normalize_level(text: str) -> str:
    m = re.search(r"(?:\bL|\bLevel\s*)?([2-4])\b", _clean(text), re.IGNORECASE)
    return f"Level {m.group(1)}" if m else _line(text)


def normalize_region(text: str) -> str:
    upper = _clean(text).upper()
    codes = [c for c in ("AP", "EU", "NA") if re.search(rf"\b{c}\b", upper)]
    return " + ".join(codes) if codes else _line(text)


def merge_regions(values: list[str]) -> str:
    codes = [c for c in ("AP", "EU", "NA") if any(re.search(rf"\b{c}\b", v.upper()) for v in values)]
    return " + ".join(codes) if codes else next((v for v in values if v), "")


def split_ids(text: str) -> tuple[str, str]:
    """SWAT ID / Bundle ID 格子 → (案例号, 行号)：bundle 号 "B/2025/00361" 做案例号，SWAT 数字号做行号。"""
    t = _clean(text)
    bundle = BUNDLE_RE.search(t)
    swat = SWAT_RE.search(t)
    case_id = bundle.group(0).upper() if bundle else (swat.group(0) if swat else t.replace("\n", " "))
    return case_id, (swat.group(0) if swat else "")


def split_part_numbers(text: str) -> list[str]:
    """一格多个零件号（换行 / 逗号 / 斜杠分隔且每段都像零件号）拆成多行；否则原样一行。"""
    t = _clean(text)
    tokens = [p.strip() for p in re.split(r"[\n,;/]+", t) if p.strip()]
    if len(tokens) > 1 and all(PN_TOKEN_RE.match(p) for p in tokens):
        return tokens
    return [t.replace("\n", " ")]


# ---------- 工作表 → 案例 ----------

def _key(header: Any) -> str:
    return re.sub(r"[^a-z0-9#]", "", str(header or "").lower())


def _column_map(row: list[Any]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, header in enumerate(row):
        k = _key(header)
        if not k:
            continue
        for field, prefixes in COLUMN_KEYS:
            if field in out:
                continue
            if any(k == p if p in EXACT_KEYS else k.startswith(p) for p in prefixes):
                out[field] = idx
                break
    return out


def _find_header(ws) -> tuple[int, dict[str, int]] | None:
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        cols = _column_map(list(row))
        if "pn" in cols and "region" in cols:
            return r, cols
    return None


def _expand_merged(ws) -> tuple[dict[tuple[int, int], Any], set[tuple[int, int]]]:
    """合并区域展开成每格都有值；返回 (值表, 合并填充的格子集合)。"""
    values: dict[tuple[int, int], Any] = {}
    filled: set[tuple[int, int]] = set()
    for rng in ws.merged_cells.ranges:
        top = ws.cell(row=rng.min_row, column=rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    values[(r, c)] = top
                    filled.add((r, c))
    return values, filled


def parse_workbook(data: bytes, filename: str, fx_settings: dict | None = None, *,
                   meeting_date: date | None = None, source_url: str = "", imported_by: str = "",
                   when: str = "") -> dict:
    """一份周文件 → {"weekNum","meetingYear","meetingDateISO","meetingDateLabel","sheet","cases":[…],"warnings":[…],"rows":n}。

    案例字典的形状与 logic.mk_case 一致，可直接入库；source 记录来源文件与行号。"""
    kw, guessed = meeting_from_name(filename)
    day = meeting_date or guessed
    if day is None:
        raise TemplateError("Cannot tell the meeting date from the file name — enter it in the import form "
                            "(expected something like 'KW39 23.09.2026_…_MM.xlsx').")
    kw, year = day.isocalendar().week, day.isocalendar().year
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws, found = None, None
    for sheet in wb.worksheets:
        found = _find_header(sheet)
        if found:
            ws = sheet
            break
    if ws is None or not found:
        raise TemplateError("This does not look like the committee template: no header row with 'PN' and 'Region' found.")
    header_row, cols = found
    merged, filled = _expand_merged(ws)

    def cell(r: int, field: str) -> Any:
        c = cols.get(field)
        return None if c is None else merged.get((r, c + 1), ws.cell(row=r, column=c + 1).value)

    rows: dict[int, dict[str, Any]] = {}
    own: dict[int, set[str]] = {}
    groups: list[list[int]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        rows[r] = {f: cell(r, f) for f in FIELDS}
        own[r] = {f for f in FIELDS if cols.get(f) is not None and (r, cols[f] + 1) not in filled and not _is_blank(rows[r][f])}
        if not own[r]:
            continue  # 空行，或只有合并填充的行
        if "num" in own[r] or not groups:
            groups.append([r])
        else:
            groups[-1].append(r)

    label = meeting_label(day)
    warnings: list[str] = []
    cases: list[dict] = []
    total_rows = 0
    for index, group in enumerate(groups, start=1):
        first = group[0]
        head = rows[first]
        if not any(own[r] & {"pn", "partDescription", "swat", "project"} for r in group):
            continue  # 没有零件也没有编号的碎行（如备注行）
        num = parse_int(head["num"]) if "num" in own[first] else None
        case_id, _ = split_ids(head["swat"])
        decision_text = _clean(head["decision"])
        decision = classify_decision(decision_text)
        case_type = normalize_type(head["type"])
        case_project = _line(head["project"])
        case_supplier = _line(head["supplier"])
        peak_total = life_total = Decimal(0)
        part_rows: list[dict] = []
        regions: list[str] = []
        for r in group:
            v = rows[r]
            region_text = normalize_region(v["region"]) if ("region" in own[r] or r == first) else ""
            if region_text:
                regions.append(region_text)
            for field in ("peak", "lifetime"):
                col = cols.get(field)
                # 自己的值逐行相加；合并填充的值整个案例只算一次（合并区域的首行）
                if col is None or _is_blank(v[field]) or ((r, col + 1) in filled and r != first):
                    continue
                amount, note = parse_amount(v[field], fx_settings)
                if note:
                    warnings.append(f"Row {r} {field} spend: {note}")
                if amount is not None:
                    if field == "peak":
                        peak_total += amount
                    else:
                        life_total += amount
            if r != first and not (own[r] & PART_FIELDS):
                continue  # 同一零件的补充行（只有备注 / 区域）
            _, row_swat = split_ids(v["swat"]) if ("swat" in own[r] or r == first) else ("", "")
            for pn in split_part_numbers(v["pn"]):
                total_rows += 1
                row: dict[str, Any] = {
                    "partNumber": pn, "partDescription": _line(v["partDescription"]),
                    "project": (_line(v["project"]) if "project" in own[r] else "") or case_project,
                    "sourcingType": (normalize_type(v["type"]) if "type" in own[r] else "") or case_type,
                    "recommendedSupplier": (_line(v["supplier"]) if "supplier" in own[r] else "") or case_supplier,
                }
                if row_swat:
                    row["sourcingId"] = row_swat
                if region_text and r != first:
                    row["region"] = region_text
                for field in ROW_MONEY:
                    amount, note = parse_amount(v[field], fx_settings, price=field in ("pcPriceCQA", "supplierPriceLanded"))
                    if note:
                        warnings.append(f"Row {r} {field}: {note}")
                    row[field] = _money_text(amount)
                for field in ROW_INT:
                    row[field] = parse_int(v[field])
                part_rows.append(row)
        if not part_rows:
            continue
        case_number = num if num is not None else index
        summary = logic.bundle_summary(part_rows)
        notes = _clean(head["notes"])
        case = logic.mk_case({
            "weekNum": kw, "caseNumber": case_number, "swatId": case_id or f"ROW-{year}-KW{kw:02d}-{case_number}",
            "partNumber": part_rows[0]["partNumber"], "partDescription": part_rows[0]["partDescription"],
            "partNumbers": part_rows if len(part_rows) > 1 else None,
            "region": merge_regions(regions), "project": summary["project"],
            "family": "", "cluster": _line(head["commodity"]),
            "parentPF": _line(head["parentPF"]), "partFamilyCode": "",
            "sourcingType": summary["sourcingType"], "recommendedSupplier": summary["recommendedSupplier"],
            "presenter": _line(head["presenter"]), "decisionLevel": normalize_level(head["decisionLevel"]),
            "peakYearSpend": _whole_eur(peak_total), "lifetimeSpend": _whole_eur(life_total), "spendCurrency": "EUR",
            "toolingPayment": "",  # 商务字段存在但为空：详情页才会显示按零件行的件价 / 模具费表
            "committeeDiscussion": notes or ("Pending for Sourcing Committee Review" if decision == "PENDING" else ""),
            "submitterComments": _clean(head["comment"]),
            "meetingDecision": decision, "decisionText": decision_text,
            "caseStatus": CASE_STATUS.get(decision, "Open"), "actionStatus": None,
            "meetingDateLabel": label, "meetingDateISO": day.isoformat(), "meetingYear": year, "createdDate": label,
            "sourcingPresentationLink": source_url if decision != "PENDING" else "",
            "followUps": [], "files": [],
            "closure": {"status": "Open", "closureDate": "—", "closureComment": "Imported from the weekly Excel.", "closureEvidence": "—"},
            "source": {"kind": SOURCE_KIND, "file": filename, "row": first, "importedAt": when, "importedBy": imported_by},
        }, [])
        logic.sync_row_commercial(case, part_rows)
        cases.append(case)
    if not cases:
        warnings.append("No case rows found under the header row.")
    return {"weekNum": kw, "meetingYear": year, "meetingDateISO": day.isoformat(), "meetingDateLabel": label,
            "sheet": ws.title, "cases": cases, "warnings": warnings, "rows": total_rows}


# ---------- 随程序一起发布的历史文件库（app/history/<年>/<KWnn, dd.mm.yyyy>/<文件>.xlsx） ----------

LIBRARY_DIR = Path(__file__).parent / "history"


def library_files() -> list[dict]:
    """库里的周文件：日期优先从文件名取，取不到用所在文件夹名（SharePoint 的 "KW46, 12.11.2025"）。"""
    out: list[dict] = []
    if not LIBRARY_DIR.is_dir():
        return out
    for path in sorted(LIBRARY_DIR.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        kw, day = meeting_from_name(path.name)
        if day is None:
            kw, day = meeting_from_name(path.parent.name)
        out.append({"path": path.relative_to(LIBRARY_DIR).as_posix(), "name": path.name, "folder": path.parent.name,
                    "weekNum": day.isocalendar().week if day else kw, "meetingYear": day.isocalendar().year if day else None,
                    "meetingDateISO": day.isoformat() if day else ""})
    return sorted(out, key=lambda f: (f["meetingYear"] or 0, f["weekNum"] or 0, f["name"]))


def library_path(rel: str) -> Path:
    """相对路径 → 库里的文件；不在库目录里、不是 .xlsx、不存在都拒绝（防目录穿越）。"""
    root = LIBRARY_DIR.resolve()
    try:
        path = (root / rel).resolve()
    except (OSError, ValueError) as exc:
        raise TemplateError("Unknown stored file.") from exc
    if root not in path.parents or not path.is_file() or path.suffix.lower() != ".xlsx":
        raise TemplateError("Unknown stored file.")
    return path


def is_imported(case: dict) -> bool:
    return (case.get("source") or {}).get("kind") == SOURCE_KIND


def previously_imported(cases: list[dict], year: int, week: int) -> list[dict]:
    """同一年同一周之前导入的案例（再次导入时被替换）。"""
    return [c for c in cases if is_imported(c) and int(c.get("meetingYear") or 0) == year and int(c.get("weekNum") or 0) == week]
