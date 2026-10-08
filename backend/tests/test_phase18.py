"""v3 Phase-18：领导第三轮反馈 —— 术语归一、演示文件库、待办自动分类与职能、按年周标签。"""

from __future__ import annotations

from copy import deepcopy

import pytest
from openpyxl import load_workbook

from app import commodities, exports, logic, migrations, service
from app.config import get_settings
from tests.test_api import VALID_SUBMISSION, _edit_payload, _get_case
from tests.test_sharepoint import submission


@pytest.fixture
def client(api_client):
    return api_client


@pytest.fixture
def admin(client):
    r = client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    assert r.status_code == 200
    return {"X-Session-Token": r.json()["token"]}


# ---------- Cluster / Commodity 归一 ----------


@pytest.mark.parametrize("raw,cluster,commodity", [
    ("CH", "Chemical", "CH Chemical"),
    ("Chemical MP Color Part", "Chemical", "CH Chemical"),
    ("CH01", "Chemical", "CH01 Functional Plastics"),
    ("Funcional Plastics", "Chemical", "CH01 Functional Plastics"),
    ("FP", "Chemical", "CH01 Functional Plastics"),
    ("Textle", "Chemical", "Textiles"),
    ("SC 03", "Metal", "SC Stamping & Mechanical Metals"),
    ("MECHANICAL METALS – NON STANDARD ( SC0300904)", "Metal", "SC Stamping & Mechanical Metals"),
    ("Stamping", "Metal", "SC Stamping & Mechanical Metals"),
    ("FF01 Forging Forming", "Metal", "FF Forging, Forming & Tubes"),
    ("Wire-Form., Tubes, Bearings & Filtration", "Metal", "FF Forging, Forming & Tubes"),
    ("TUBES", "Metal", "FF Forging, Forming & Tubes"),
    ("CA04", "Metal", "CA Casting"),
    ("Powder Metal", "Metal", "Powder Metal"),
    ("EL 02", "Electronics, EA and Pyro", "EL Electronics"),
    ("Electric Mechanicals", "Electronics, EA and Pyro", "EA Electromechanical"),
    ("EA03011", "Electronics, EA and Pyro", "EA Electromechanical"),
    ("Pyro MGG", "Electronics, EA and Pyro", "Pyro"),
    ("Metal", "Metal", ""),  # 在线登记：已经是 Cluster
])
def test_commodity_text_maps_to_cluster_and_commodity(raw, cluster, commodity):
    assert commodities.classify(raw) == {"cluster": cluster, "commodity": commodity, "known": True}


@pytest.mark.parametrize("raw,kept", [("MI 01", "MI01"), ("SO02", "SO02"), ("Yajun Zhao", "Yajun Zhao")])
def test_unknown_commodity_is_kept_as_written(raw, kept):
    assert commodities.classify(raw) == {"cluster": kept, "commodity": kept, "known": False}
    assert commodities.classify("") == {"cluster": "", "commodity": "", "known": False}


def test_normalize_commodities_only_touches_imported_or_foreign_clusters():
    online = logic.mk_case({"swatId": "S1", "cluster": "Metal"}, [])
    imported = logic.mk_case({"swatId": "S2", "cluster": "SC06", "source": {"kind": "excel"}}, [])
    foreign = logic.mk_case({"swatId": "S3", "cluster": "Functional plastics"}, [])
    changed = service.normalize_commodities([online, imported, foreign])
    assert [c["swatId"] for c in changed] == ["S2", "S3"]
    assert online == logic.mk_case({"swatId": "S1", "cluster": "Metal"}, []) and "commodityRaw" not in online
    assert (imported["cluster"], imported["commodity"], imported["commodityRaw"]) == ("Metal", "SC Stamping & Mechanical Metals", "SC06")
    assert (foreign["cluster"], foreign["commodity"], foreign["commodityRaw"]) == ("Chemical", "CH01 Functional Plastics", "Functional plastics")
    # 再算一次：原文保留，结果稳定
    assert service.normalize_commodities([imported, foreign]) == []


async def test_startup_migration_normalizes_once_per_vocabulary_version(isolated_app):
    from app import db

    async with db.SessionLocal() as session:
        await service.init_db(session)
        await service.insert_cases(session, [logic.mk_case({"swatId": "X1", "cluster": "EL 02", "source": {"kind": "excel"}}, [])])
        await service._set_kv(session, service.COMMODITY_KEY, {"version": commodities.VERSION - 1})
        await session.commit()
        await service._migrate_commodities(session)
        await session.commit()
        case = next(c for c in (await service.load_state(session))["cases"] if c["swatId"] == "X1")
        assert (case["cluster"], case["commodity"], case["commodityRaw"]) == ("Electronics, EA and Pyro", "EL Electronics", "EL02")
        assert (await service._kv_doc(session, service.COMMODITY_KEY)) == {"version": commodities.VERSION}


async def test_normalize_clusters_command_previews_then_applies(isolated_app, capsys):
    from app import db

    async with db.SessionLocal() as session:
        await service.init_db(session)
        await service.insert_cases(session, [
            logic.mk_case({"swatId": "Y1", "cluster": "SC Stamping", "source": {"kind": "excel"}}, []),
            logic.mk_case({"swatId": "Y2", "cluster": "MI01", "source": {"kind": "excel"}}, []),
        ])
        await session.commit()
    assert await migrations.run(["normalize-clusters"], db.SessionLocal) == 0
    out = capsys.readouterr().out
    assert "SC Stamping" in out and "-> Metal" in out and "not in the vocabulary (kept as written): MI01 (1)" in out
    assert "dry run: 2 of" in out
    async with db.SessionLocal() as session:
        cases = (await service.load_state(session))["cases"]
        assert next(c for c in cases if c["swatId"] == "Y1")["cluster"] == "SC Stamping"  # 预览不写库
    assert await migrations.run(["normalize-clusters", "--apply"], db.SessionLocal) == 0
    assert "applied: 2 of" in capsys.readouterr().out
    async with db.SessionLocal() as session:
        cases = (await service.load_state(session))["cases"]
        assert next(c for c in cases if c["swatId"] == "Y1")["cluster"] == "Metal"
        assert next(c for c in cases if c["swatId"] == "Y2")["cluster"] == "MI01"


def test_edit_keeps_an_unmapped_cluster_text(client, admin):
    """编辑页的 Cluster 放宽为文本：词表认不出的 "MI01" 保存时原样保留（之前的 Literal 会 422）。"""
    assert client.post("/api/submissions", json=VALID_SUBMISSION).status_code == 200
    assert client.post("/api/submissions/SUB-0002/confirm", headers=admin).status_code == 200
    case = _get_case(client, "C0023")
    r = client.put(f"/api/cases/{case['id']}", headers=admin, json=_edit_payload(case, cluster="MI01"))
    assert r.status_code == 200 and r.json()["case"]["cluster"] == "MI01"
    r = client.put(f"/api/cases/{case['id']}", headers=admin, json=_edit_payload(case, cluster="x" * 121))
    assert r.status_code == 422


# ---------- 待办：自由文本标签 → 关键词分类；负责人职能 ----------


@pytest.mark.parametrize("task,tags,notes,expected", [
    ("CARRIE TEST SAMPLES", [], "", ["CQA", "Technical"]),  # samples → CQA，test → Technical
    ("Confirm SOP timing with supplier", [], "", ["Timing"]),
    ("Negotiate price reduction and BPG", ["Quality"], "", ["CQA", "BPG", "Saving"]),
    ("Prepare dual source strategy for KW45", [], "VAVE workshop ideas", ["Timing", "Supplier strategy", "Further VAVE"]),
    ("Check PPAP samples", ["Timing"], "", ["CQA", "Timing"]),
    ("Backward compatibility review", [], "", []),  # "kw" 藏在 backward 里不算
    ("Review drawing and test plan", ["Cost"], "", ["Technical", "Saving"]),
    ("", ["Compliance / Documentation", "Capacity"], "", ["CQA", "Volume"]),  # 模板旧标签按同义词归类
    ("Follow up", [], "", []),
])
def test_tasks_are_categorized_by_keywords(task, tags, notes, expected):
    assert logic.categorize_task({"task": task, "tags": tags, "notes": notes}) == expected


def test_task_categories_fall_back_to_rules_for_old_records():
    assert logic.task_categories({"task": "x", "tags": ["Saving"], "categories": ["BPG"]}) == ["BPG"]  # 保存过的以保存为准
    assert logic.task_categories({"task": "Volume ramp-up", "tags": []}) == ["Volume"]
    rules = logic.task_rules()
    assert rules["categories"] == logic.ACTION_CATEGORIES and set(rules["keywords"]) == set(logic.ACTION_CATEGORIES)
    assert rules["tagSynonyms"]["quality"] == "CQA" and "Purchasing" in rules["functions"]


def test_edit_saves_owner_function_and_derived_categories(client, admin):
    assert client.post("/api/submissions", json=VALID_SUBMISSION).status_code == 200
    assert client.post("/api/submissions/SUB-0002/confirm", headers=admin).status_code == 200
    case = _get_case(client, "C0023")
    task = {"id": "FU-T1", "task": "Send PPAP samples", "responsible": "Carrie Wang", "ownerFunction": "Supplier Quality",
            "dueDate": "2027-10-01", "status": "Open", "notes": "", "tags": ["Quality", "Timing"], "notifyEmail": "carrie.wang@zf.com"}
    r = client.put(f"/api/cases/{case['id']}", headers=admin, json=_edit_payload(
        case, meetingDecision="APPROVED", sourcingPresentationLink="https://zf.sharepoint.com/x.pptx", followUps=[task]))
    assert r.status_code == 200, r.text
    saved = r.json()["case"]["followUps"][0]
    assert saved["ownerFunction"] == "Supplier Quality" and saved["tags"] == ["Quality", "Timing"]
    assert saved["categories"] == ["CQA", "Timing"]  # 由关键词与标签同义词判定，不是手点的
    assert client.get("/api/bootstrap").json()["taskRules"]["categories"] == logic.ACTION_CATEGORIES
    # 职能超长拒绝；省略时为空
    r = client.put(f"/api/cases/{case['id']}", headers=admin, json=_edit_payload(
        case, meetingDecision="APPROVED", sourcingPresentationLink="https://zf.sharepoint.com/x.pptx",
        followUps=[dict(task, ownerFunction="f" * 81)]))
    assert r.status_code == 422
    plain = {k: v for k, v in task.items() if k != "ownerFunction"}
    r = client.put(f"/api/cases/{case['id']}", headers=admin, json=_edit_payload(
        case, meetingDecision="APPROVED", sourcingPresentationLink="https://zf.sharepoint.com/x.pptx", followUps=[plain]))
    assert r.status_code == 200 and r.json()["case"]["followUps"][0]["ownerFunction"] == ""


# ---------- 导出：周标签带年份；Actions 表多了 Tags / Function ----------


def _sheet(data: bytes, title: str):
    return load_workbook(__import__("io").BytesIO(data))[title]


def test_week_label_uses_the_case_year():
    assert logic.wk(51, 2025) == "2025-KW51" and logic.wk(51) == "2026-KW51"
    assert logic.week_label({"weekNum": 51, "meetingYear": 2025}) == "2025-KW51"
    assert logic.week_label({"weekNum": 51, "meetingDateISO": "2025-12-17"}) == "2025-KW51"
    assert logic.week_label({"weekNum": 3, "meetingDateLabel": "Wed, Jan 15, 2025"}) == "2025-KW03"
    assert logic.week_label({"weekNum": 7}) == "2026-KW07"  # 模板演示数据没有年份：按当年
    old = logic.mk_case({"swatId": "OLD", "weekNum": 51, "meetingYear": 2025, "meetingDateISO": "2025-12-17"}, [])
    ws = _sheet(exports.build("agenda", [old], "2025-KW51"), "Agenda")
    headers = [c.value for c in ws[4]]
    assert ws.cell(row=5, column=headers.index("KW") + 1).value == "2025-KW51"


def test_actions_sheet_lists_category_tags_and_function():
    task = {"id": "T1", "task": "Negotiate consignment", "tags": ["Supplier strategy", "Saving"], "ownerFunction": "Purchasing",
            "responsible": "SP ABC", "dueDate": "2027-02-18", "status": "Open", "notifyEmail": "sp@example.com"}
    case = logic.mk_case({"swatId": "SWAT-1", "meetingDecision": "APPROVED", "followUps": [task]}, [])
    ws = _sheet(exports.build("minutes", [case], "2026-KW39"), "Actions")
    headers = [c.value for c in ws[4]]
    assert headers[4:10] == ["Task", "Category", "Tags", "Owner", "Function", "Due Date"]
    row = dict(zip(headers, [c.value for c in ws[5]], strict=True))
    assert (row["Category"], row["Tags"], row["Function"]) == ("Supplier strategy, Saving", "Supplier strategy, Saving", "Purchasing")


def test_week_export_label_follows_the_requested_year(client):
    r = client.get("/api/exports/agenda", params={"week": 51, "year": 2025})
    assert r.status_code == 200 and 'filename="Sourcing_Committee_Agenda_2025-KW51.xlsx"' in r.headers["content-disposition"]


# ---------- 演示文件库：从服务器上选文件挂到案例 ----------


@pytest.fixture
def library(tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "upload_dir", str(tmp_path / "uploads"))
    root = tmp_path / "presentations"
    (root / "KW39, 23.09.2026").mkdir(parents=True)
    (root / "KW39, 23.09.2026" / "MBEAL deck.pptx").write_bytes(b"deck-bytes")
    (root / "notes.txt").write_text("not a presentation")
    (root / "~$lock.pptx").write_bytes(b"lock")
    (root / "Final report.pdf").write_bytes(b"pdf")
    (root / "empty.pptx").write_bytes(b"")
    (root / "KW39 MM.xlsx").write_bytes(b"xlsx")  # 周会 Excel 不进 Presentations
    monkeypatch.setattr(get_settings(), "presentation_library_dir", str(root))
    return root


async def _manager(client):
    r = await client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    return {"X-Session-Token": r.json()["token"]}


async def _case_with_files(client, headers, number=40011):
    r = await client.post("/api/submissions", json=submission(number))
    assert r.status_code == 200, r.text
    return (await client.post(f"/api/submissions/{r.json()['subId']}/confirm", headers=headers)).json()["case"]


async def test_library_lists_presentation_files_for_managers_only(async_client, library):
    assert (await async_client.get("/api/presentation-library")).status_code == 403
    r = await async_client.get("/api/presentation-library", headers=await _manager(async_client))
    assert r.status_code == 200
    listed = r.json()["files"]
    assert [(f["folder"], f["name"], f["size"]) for f in listed] == [
        ("", "empty.pptx", 0), ("", "Final report.pdf", 3), ("KW39, 23.09.2026", "MBEAL deck.pptx", 10)]
    assert listed[2]["path"] == "KW39, 23.09.2026/MBEAL deck.pptx" and listed[2]["modified"]
    assert r.json()["folder"] == str(library)
    get_settings().presentation_library_dir = str(library / "missing")
    assert (await async_client.get("/api/presentation-library", headers=await _manager(async_client))).json()["files"] == []


async def test_attach_from_library_copies_and_renames_and_keeps_the_original(async_client, library):
    from app import files

    headers = await _manager(async_client)
    case = await _case_with_files(async_client, headers)
    assert (await async_client.post(f"/api/cases/{case['id']}/files/from-library", json={"path": "KW39, 23.09.2026/MBEAL deck.pptx"})).status_code == 403
    r = await async_client.post(f"/api/cases/{case['id']}/files/from-library", headers=headers, json={"path": "KW39, 23.09.2026/MBEAL deck.pptx"})
    assert r.status_code == 200, r.text
    rec = r.json()["file"]
    assert rec["name"] == "2026.09.24 Test Project Test Part 1, Test Supplier.pptx" and rec["originalName"] == "MBEAL deck.pptx"
    assert rec["size"] == 10 and rec["store"] == "local" and rec["uploadedBy"] == "Sourcing admin"
    assert (await async_client.get(f"/api/files/{rec['id']}")).content == b"deck-bytes"
    assert (library / "KW39, 23.09.2026" / "MBEAL deck.pptx").read_bytes() == b"deck-bytes"  # 库里的原件不动
    assert not list((files.upload_root() / ".incoming").iterdir())
    assert [f["id"] for f in r.json()["case"]["files"]] == [rec["id"]]
    # 同一份可以挂到第二个案例
    other = await _case_with_files(async_client, headers, 40012)
    assert (await async_client.post(f"/api/cases/{other['id']}/files/from-library", headers=headers, json={"path": "KW39, 23.09.2026/MBEAL deck.pptx"})).status_code == 200


@pytest.mark.parametrize("path,status", [
    ("../presentations/Final report.pdf", 404),  # 解析后仍在库里也拒绝（必须是库内相对路径）
    ("../../etc/passwd", 404), ("notes.txt", 404), ("~$lock.pptx", 404), ("nope.pptx", 404), ("", 422), ("empty.pptx", 422), ("KW39 MM.xlsx", 404),
])
async def test_attach_rejects_paths_outside_the_library_or_wrong_files(async_client, library, path, status):
    headers = await _manager(async_client)
    case = await _case_with_files(async_client, headers, 40013)
    r = await async_client.post(f"/api/cases/{case['id']}/files/from-library", headers=headers, json={"path": path})
    assert r.status_code == status, r.text


async def test_attach_respects_the_per_case_file_limit(async_client, library):
    from app import files

    headers = await _manager(async_client)
    case = await _case_with_files(async_client, headers, 40014)
    for _ in range(files.MAX_FILES_PER_RECORD):
        assert (await async_client.post(f"/api/cases/{case['id']}/files/from-library", headers=headers, json={"path": "Final report.pdf"})).status_code == 200
    r = await async_client.post(f"/api/cases/{case['id']}/files/from-library", headers=headers, json={"path": "Final report.pdf"})
    assert r.status_code == 422 and "at most" in r.json()["detail"]
    assert (await async_client.post("/api/cases/C9999/files/from-library", headers=headers, json={"path": "Final report.pdf"})).status_code == 404


def test_imported_cases_carry_cluster_commodity_and_raw(client, admin):
    """导入接口：Commodity 原文归一，认不出的进预览警告。"""
    from tests.test_excel_import import SAMPLE_MERGES, SAMPLE_ROWS, _workbook

    rows = deepcopy(SAMPLE_ROWS)
    r = client.post("/api/import/excel", params={"name": "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", "commit": "false"},
                    content=_workbook(rows, SAMPLE_MERGES), headers={**admin, "Content-Type": "application/octet-stream"})
    assert r.status_code == 200, r.text
    assert r.json()["cases"] == 4 and not [w for w in r.json()["warnings"] if "vocabulary" in w]
    rows[0]["Commodity"] = "Yajun Zhao"
    r = client.post("/api/import/excel", params={"name": "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", "commit": "true"},
                    content=_workbook(rows, SAMPLE_MERGES), headers={**admin, "Content-Type": "application/octet-stream"})
    assert r.status_code == 200, r.text
    assert [w for w in r.json()["warnings"] if "vocabulary" in w] == ["Row 2: commodity 'Yajun Zhao' is not in the vocabulary yet — kept as written."]
    imported = [c for c in r.json()["snapshot"]["cases"] if c.get("source")]
    assert (imported[0]["cluster"], imported[0]["commodity"], imported[0]["commodityRaw"]) == ("Yajun Zhao", "Yajun Zhao", "Yajun Zhao")
    assert (imported[1]["cluster"], imported[1]["commodity"], imported[1]["commodityRaw"]) == ("Electronics, EA and Pyro", "EA Electromechanical", "EA05")
