"""全项目检查回归：跨年周、金额合法性与导出后的币种往返。"""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app import excel_import, exports, files, logic, presentation_match
from app.config import get_settings

FX = {"perEur": {"USD": "1.25", "CNY": "8"}}


@pytest.mark.parametrize("name, expected", [
    ("KW53 01.01.2021_MM.xlsx", (53, date(2020, 12, 30))),
    ("KW01 31.12.2025_MM.xlsx", (1, date(2025, 12, 31))),
    ("KW53 2025_MM.xlsx", (53, None)),
    ("KW00 2026_MM.xlsx", (0, None)),
    ("KW99 23.09.2026_MM.xlsx", (99, None)),
])
def test_meeting_name_handles_cross_year_and_invalid_iso_weeks(name, expected):
    assert excel_import.meeting_from_name(name) == expected


def test_presentation_library_can_match_a_cross_year_week_and_tolerate_bad_folders():
    case = {"id": "ROW-1", "swatId": "110000000123", "meetingYear": 2020, "weekNum": 53, "caseNumber": 1}
    result = presentation_match.match("KW53, 01.01.2021/1. deck.pdf", [case])
    assert result["rowId"] == "ROW-1" and result["by"] == "week + case #"
    assert presentation_match.match("KW99 2026/deck.pdf", [case]) is None


@pytest.mark.parametrize("value", ["-100 EUR", "-2k EUR", "-1.25 USD", "-100", "-2k", -100, -0.5,
                                  "−100 EUR", "- 100 EUR", "− 1.25 USD", "−2k"])
def test_import_rejects_negative_amounts_instead_of_turning_them_positive(value):
    amount, warning = excel_import.parse_amount(value, FX)
    assert amount is None and "negative" in warning.lower()


@pytest.mark.parametrize("value, expected", [
    ("1,76 -2,17 M€", 2_170_000),  # KW12 / KW28 历史文件的原文：区间取紧挨单位的那个数
    ("1,76 - 2,17 Mio EUR", 2_170_000),
    ("1,76-2,17M€", 2_170_000),
    ("1,76 – 2,17 M€", 2_170_000),
])
def test_dash_between_numbers_is_a_range_not_a_minus_sign(value, expected):
    assert excel_import.parse_amount(value, FX) == (expected, None)


@pytest.mark.parametrize("value", ["5-6", "5 - 6k"])
def test_bare_range_stays_unreadable_instead_of_merging_digits(value):
    amount, warning = excel_import.parse_amount(value, FX)
    assert amount is None and "cannot read amount" in warning and value in warning


@pytest.mark.parametrize("currency, expected", [
    ("RMB", 10), ("rmb", 10), ("¥", 10), ("CNY", 10),
    ("€", 80), ("Euro", 80), (" eur ", 80), ("", 80),
    ("$", 64), ("usd", 64),
])
def test_currency_column_accepts_common_spellings(currency, expected):
    amount, _ = excel_import.parse_amount(80, FX, currency=currency)
    assert amount == expected


def test_unknown_currency_column_warns_and_skips():
    amount, warning = excel_import.parse_amount(80, FX, currency="GBP")
    assert amount is None and "unsupported currency 'GBP'" in warning


@pytest.mark.parametrize("value", [float("nan"), float("inf"), -float("inf")])
def test_import_rejects_nonfinite_numbers(value):
    amount, warning = excel_import.parse_amount(value, FX)
    assert amount is None and warning


def _exported_case(currency):
    return logic.mk_case({
        "weekNum": 39, "meetingYear": 2026, "meetingDateISO": "2026-09-23", "caseNumber": 1,
        "swatId": "110000000123", "region": "EU", "partNumber": "00012345", "partDescription": "Bracket",
        "spendCurrency": currency, "peakYearSpend": 1000, "lifetimeSpend": 10000,
        "pcPriceCQA": "10", "supplierPriceLanded": "8", "toolingCQA": "100", "supplierToolingCost": "80",
        "averageVolume": 100, "lifetimeVolume": 1000, "meetingDecision": "PENDING",
    }, [])


@pytest.mark.parametrize("currency, rate", [("EUR", Decimal(1)), ("USD", Decimal("1.25")), ("CNY", Decimal(8))])
def test_exported_numeric_amounts_reimport_in_the_declared_currency(currency, rate):
    original = _exported_case(currency)
    data = exports.build_minutes([original], "2026-KW39")
    report = excel_import.parse_workbook(data, exports.filename("minutes", [original]), FX)
    imported, = report["cases"]
    assert imported["spendCurrency"] == "EUR"
    assert imported["peakYearSpend"] == int(Decimal(1000) / rate)
    assert imported["lifetimeSpend"] == int(Decimal(10000) / rate)
    for field in excel_import.ROW_MONEY:
        assert Decimal(imported[field]) == Decimal(original[field]) / rate
    assert imported["averageVolume"] == 100 and imported["lifetimeVolume"] == 1000
    assert imported["partNumber"] == "00012345"


def test_foreign_currency_column_without_rate_warns_and_does_not_relabel_amounts():
    original = _exported_case("USD")
    data = exports.build_minutes([original], "2026-KW39")
    report = excel_import.parse_workbook(data, exports.filename("minutes", [original]), {"perEur": {}})
    imported, = report["cases"]
    assert imported["peakYearSpend"] == 0 and imported["lifetimeSpend"] == 0
    assert imported["pcPriceCQA"] is None
    assert any("no USD rate" in warning for warning in report["warnings"])


@pytest.mark.parametrize("value", ["20", "20k", 20])
def test_amount_without_inline_unit_obeys_currency_column(value):
    amount, warning = excel_import.parse_amount(value, FX, currency="USD")
    assert amount == (16000 if value == "20k" else 16)
    assert "USD converted" in warning
    assert excel_import.parse_amount("20 EUR", FX, currency="USD") == (20, None)


def test_invalid_week_in_filename_does_not_block_explicit_meeting_date():
    data = exports.build_minutes([_exported_case("EUR")], "2026-KW39")
    report = excel_import.parse_workbook(data, "KW53 2025_MM.xlsx", FX, meeting_date=date(2026, 9, 23))
    assert report["meetingDateISO"] == "2026-09-23"


def test_bad_week_folder_does_not_break_presentation_library_endpoint(api_client, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "presentation_library_dir", str(tmp_path / "presentations"))
    folder = files.library_root() / "KW99 2026"
    folder.mkdir(parents=True)
    (folder / "deck.pdf").write_bytes(b"presentation")
    token = api_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"}).json()["token"]
    response = api_client.get("/api/presentation-library", headers={"X-Session-Token": token})
    assert response.status_code == 200
    assert response.json()["files"][0]["match"] is None


@pytest.mark.parametrize("kind", ["agenda", "minutes"])
def test_export_ignores_xml_control_characters_and_keeps_text_formula_safe(kind):
    case = _exported_case("EUR")
    case["submitterComments"] = "\x01=1+1\x0b\ufffe\uffff"
    case["partDescription"] = "Line 1\nLine 2\tTabs\rReturn\x00\ud800"
    data = exports.build(kind, [case], "2026-KW39")
    ws = load_workbook(io.BytesIO(data)).active
    headers = {cell.value: cell.column for cell in ws[exports.HEADER_ROW]}
    comment = ws.cell(exports.HEADER_ROW + 1, headers["Comment"])
    assert comment.value == "=1+1" and comment.data_type == "s"
    # XML 解析器把单独的回车规范化为换行；两种换行都仍保留。
    assert ws.cell(exports.HEADER_ROW + 1, headers["Part Description"]).value == "Line 1\nLine 2\tTabs\nReturn"
    assert case["submitterComments"] == "\x01=1+1\x0b\ufffe\uffff"
