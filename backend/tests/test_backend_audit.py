"""项目巡检回归：账号并发、请求边界与跨年会议；数据和外发通道均由 conftest 隔离。"""

import asyncio
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app import accounts, logic, main, service
from app.schemas import UserUpdateIn
from tests.test_api import VALID_SUBMISSION


async def test_login_does_not_overwrite_a_queued_account_update(async_client, monkeypatch):
    original = service.get_user
    loaded, release = asyncio.Event(), asyncio.Event()
    first = True

    async def delayed_read(session, email):
        nonlocal first
        user = await original(session, email)
        if email == "test.user@zf.com" and first:
            first = False
            loaded.set()
            await release.wait()
        return user

    monkeypatch.setattr(service, "get_user", delayed_read)
    login = asyncio.create_task(async_client.post("/api/auth/login", json={"email": "test.user@zf.com"}))
    await asyncio.wait_for(loaded.wait(), timeout=3)
    # 直接调用已通过管理员鉴权的处理函数，让账号更新排在当前读取之后等待同一把写锁。
    update = asyncio.create_task(main.update_account(
        "test.user@zf.com", UserUpdateIn(name="Updated Person", role="npi_manager", disabled=True),
    ))
    await asyncio.sleep(0)
    release.set()
    response, _ = await asyncio.gather(login, update)
    assert response.status_code == 200
    async with main.SessionLocal() as session:
        stored = await original(session, "test.user@zf.com")
        assert (stored["name"], stored["role"], stored["disabled"]) == ("Updated Person", "npi_manager", True)
        assert await service.session_user(session, response.json()["token"]) is None


async def test_password_change_does_not_undo_a_queued_password_reset(async_client, monkeypatch):
    login = await async_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    token = login.json()["token"]
    original = service.session_user
    loaded, release = asyncio.Event(), asyncio.Event()
    first = True

    async def delayed_read(session, current_token):
        nonlocal first
        user = await original(session, current_token)
        if current_token == token and first:
            first = False
            loaded.set()
            await release.wait()
        return user

    monkeypatch.setattr(service, "session_user", delayed_read)
    change = asyncio.create_task(async_client.post(
        "/api/auth/password", headers={"X-Session-Token": token},
        json={"currentPassword": "test-admin", "newPassword": "my-new-password"},
    ))
    await asyncio.wait_for(loaded.wait(), timeout=3)
    reset = asyncio.create_task(main.reset_account_password("admin@zf.com"))
    await asyncio.sleep(0)
    release.set()
    response, result = await asyncio.gather(change, reset)
    assert response.status_code in (200, 401)
    async with main.SessionLocal() as session:
        stored = await service.get_user(session, "admin@zf.com")
        assert accounts.verify_password(result["temporaryPassword"], stored["passwordHash"])
        assert stored["mustChangePassword"] is True
        assert await original(session, token) is None


async def test_concurrent_failed_logins_obey_the_lockout(async_client):
    responses = await asyncio.gather(*[
        async_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "incorrect"})
        for _ in range(main.LOGIN_MAX_FAILS + 3)
    ])
    assert sorted(response.status_code for response in responses) == [401] * main.LOGIN_MAX_FAILS + [429] * 3


@pytest.mark.parametrize("field", ["name", "role", "disabled"])
def test_account_update_rejects_explicit_null(api_client, field):
    login = api_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    headers = {"X-Session-Token": login.json()["token"]}
    before = api_client.get("/api/users", headers=headers).json()["users"]
    response = api_client.put("/api/users/test.user@zf.com", json={field: None}, headers=headers)
    assert response.status_code == 422
    assert api_client.get("/api/users", headers=headers).json()["users"] == before


async def test_bootstrap_admin_restore_revokes_old_sessions(async_client):
    login = await async_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    token = login.json()["token"]
    async with main.SessionLocal() as session:
        admin = await service.get_user(session, "admin@zf.com")
        # 模拟运维恢复前所有管理员均不可用的数据库状态。
        await service.save_user(session, {**admin, "disabled": True})
        await session.commit()
        await service.init_db(session)
        assert (await service.get_user(session, "admin@zf.com"))["disabled"] is False
        assert await service.session_user(session, token) is None


def test_temporary_admin_cannot_upload_to_another_submission(api_client, monkeypatch):
    login = api_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    admin = {"X-Session-Token": login.json()["token"]}
    sub = api_client.post("/api/submissions", json=VALID_SUBMISSION).json()
    api_client.post("/api/auth/register", json={"email": "temporary@zf.com", "name": "Temporary Admin"})
    promoted = api_client.put("/api/users/temporary@zf.com", json={"role": "admin"}, headers=admin).json()
    temporary_login = api_client.post("/api/auth/login", json={
        "email": "temporary@zf.com", "password": promoted["temporaryPassword"],
    })

    async def unexpected_receive(request):
        raise AssertionError("A temporary-password account must be rejected before receiving or storing a file")

    monkeypatch.setattr(main.files, "receive", unexpected_receive)
    response = api_client.post(
        f"/api/submissions/{sub['subId']}/files?name=evidence.pdf", content=b"PDF",
        headers={"X-Session-Token": temporary_login.json()["token"]},
    )
    assert response.status_code == 403
    assert "temporary password" in response.json()["detail"]


def test_confirmation_uses_iso_year_at_year_boundary(api_client, monkeypatch):
    monkeypatch.setattr(logic, "na_now", lambda: datetime.fromisoformat("2025-12-20T09:00:00").replace(tzinfo=logic.app_tz()))
    login = api_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    admin = {"X-Session-Token": login.json()["token"]}
    for number in ("SWAT-YEAR-END-1", "SWAT-YEAR-END-2"):
        submitted = api_client.post("/api/submissions", json={
            **VALID_SUBMISSION, "caseNo": number, "meetingDateISO": "2025-12-31",
        })
        assert submitted.status_code == 200, submitted.text
        confirmed = api_client.post(f"/api/submissions/{submitted.json()['subId']}/confirm", headers=admin)
        case = confirmed.json()["case"]
        assert (case["meetingYear"], case["weekNum"]) == (2026, 1)
    assert case["caseNumber"] == 2
    exported = api_client.get("/api/exports/agenda", params={"week": 1, "year": 2026})
    sheet = load_workbook(BytesIO(exported.content)).active
    values = {str(value) for row in sheet.iter_rows(values_only=True) for value in row}
    assert {"SWAT-YEAR-END-1", "SWAT-YEAR-END-2"} <= values


def test_default_week_label_uses_iso_year(monkeypatch):
    monkeypatch.setattr(logic, "na_now", lambda: datetime.fromisoformat("2025-12-31T09:00:00").replace(tzinfo=logic.app_tz()))
    assert logic.wk(1) == "2026-KW01"


def test_export_rejects_a_week_that_does_not_exist_in_the_year(api_client):
    response = api_client.get("/api/exports/agenda", params={"week": 53, "year": 2025})
    assert response.status_code == 422
    assert "ISO week" in response.json()["detail"]
    assert api_client.get("/api/exports/agenda", params={"week": 53, "year": 2026}).status_code == 200


def test_registration_limit_prunes_inactive_addresses(monkeypatch):
    monkeypatch.setattr(main.time, "time", lambda: 10000)
    monkeypatch.setattr(main, "_register_times", {"inactive": [1.0], "empty": [], "active": [9999.0]})
    assert main._register_allowed("new-address") is True
    assert main._register_times == {"active": [9999.0]}
