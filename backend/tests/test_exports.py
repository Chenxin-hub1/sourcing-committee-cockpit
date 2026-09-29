"""v3 Phase-4：议程 / 纪要 .xlsx 导出，版式按反馈人的 MM 模板。读回生成的文件核对内容，不依赖 Excel。"""

import io
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from app import exports, logic

WHEN = datetime(2026, 9, 24, 9, 0, tzinfo=ZoneInfo("America/New_York"))
ROW_A = {"partNumber": "R003H147A", "partDescription": "ECU Housing", "peakYearSpend": 1104000, "lifetimeSpend": 3220000,
         "pcPriceCQA": "9.66", "supplierPriceLanded": "9.384", "averageVolume": 378979, "lifetimeVolume": 3410814,
         "toolingCQA": "92000", "supplierToolingCost": "0"}
ROW_B = {"partNumber": "R004D369A", "partDescription": "ECU Cover", "peakYearSpend": 139000, "lifetimeSpend": 852000,
         "pcPriceCQA": "0.12", "supplierPriceLanded": "0.18", "averageVolume": 525895, "lifetimeVolume": 4733055,
         "toolingCQA": "0", "supplierToolingCost": "49529"}


def _case(**over):
    base = logic.mk_case({"weekNum": 39, "caseNumber": 2, "swatId": "SWAT-77001", "region": "EU + NA", "project": "MBEAL",
                          "cluster": "Metal", "parentPF": "CH01016", "sourcingType": "New", "decisionLevel": "Level 2",
                          "recommendedSupplier": "Stara Plastics", "presenter": "E2E", "meetingDateLabel": "Wed, Sep 30, 2026",
                          "spendCurrency": "EUR", "toolingPayment": "MPC", "bpgAvailable": "Yes", "fraAvailable": "Yes",
                          "submitterComments": "=SUM(A1)", "meetingDecision": "CONDITIONAL APPROVAL", "actionStatus": "Open",
                          "committeeDiscussion": "Approved subject to consignment.",
                          "followUps": [{"id": "FU-1", "task": "Negotiate consignment", "responsible": "SP ABC", "dueDate": "2027-02-18",
                                         "status": "Open", "tags": ["Supplier strategy", "Saving"], "notifyEmail": "sp@example.com", "cc": "", "notes": ""},
                                        {"id": "FU-2", "task": "Send BP", "responsible": "SP ABC", "dueDate": "2026-10-01", "status": "Closed", "tags": ["BPG"]}],
                          **{k: v for k, v in ROW_A.items() if k not in ("partNumber", "partDescription")},
                          "partNumber": ROW_A["partNumber"], "partDescription": ROW_A["partDescription"]}, [])
    base.update(over)
    return base


def _multi(**over):
    return _case(**{"swatId": "SWAT-77002", "caseNumber": 3, "partNumbers": [ROW_A, ROW_B],
                    "peakYearSpend": 1243000, "lifetimeSpend": 4072000, "pcPriceCQA": None, "supplierPriceLanded": None, **over})


def _sheet(data: bytes, name: str):
    return load_workbook(io.BytesIO(data))[name]


def _rows(ws, header_row=4):
    headers = [c.value for c in ws[header_row]]
    # openpyxl 把空单元格读回 None；这里统一成 ""，让断言只关心业务值
    return headers, [dict(zip(headers, ("" if c.value is None else c.value for c in row), strict=True)) for row in ws.iter_rows(min_row=header_row + 1)]


def test_agenda_follows_the_mm_template_one_row_per_part_number():
    data = exports.build("agenda", [_multi(), _case(), _case(caseNumber=1, swatId="SWAT-1", meetingDecision="PENDING")], "2026-KW39", WHEN)
    ws = _sheet(data, "Agenda")
    assert ws["A1"].value == "Sourcing Committee Agenda — 2026-KW39" and ws["A1"].font.color.rgb.endswith("1F4E79")
    assert "Sep 24" in ws["A2"].value and "negative value means below CQA" in ws["A2"].value
    headers, rows = _rows(ws)
    assert headers[:11] == ["#", "KW", "Meeting Date", "Region", "Type", "Project / Product", "Commodity", "Parent PF Code",
                            "SWAT ID / Bundle ID", "Part Description", "PN"]
    assert "Sourcing Decision" not in headers and "Notes / Task / Resp. / Due Date" not in headers  # 议程不带会后信息
    # 按案例号排序；多零件案例占两行，案例级单元格合并、只在第一行有值
    assert [r["PN"] for r in rows] == ["R003H147A", "R003H147A", "R003H147A", "R004D369A"]
    assert [r["SWAT ID / Bundle ID"] for r in rows] == ["SWAT-1", "SWAT-77001", "SWAT-77002", ""]
    assert "I7:I8" in {str(m) for m in ws.merged_cells.ranges}
    assert rows[2]["Region"] == "EU, NA" and rows[2]["Decision Level"] == "L2" and rows[2]["Currency"] == "EUR"
    # v3 Phase-15：Peak Year / LT Spend 按整个 bundle 在案例级（合并单元格，只在第一行有值）
    assert rows[2]["Peak Year Spend"] == 1243000 and rows[2]["LT Spend"] == 4072000 and rows[3]["Peak Year Spend"] == ""
    assert float(rows[3]["CQA PC Price"]) == 0.12 and rows[3]["Average Volume"] == 525895
    # 项目 / 类型 / 供应商按零件行；旧多零件案例的行上没有 → 回退到案例级值
    assert [r["Project / Product"] for r in rows[2:]] == ["MBEAL", "MBEAL"] and rows[3]["Recommended Supplier(s)"] == "Stara Plastics"
    # 金额列带案例币种的符号（模板节省额的写法 #,##0 [$€-407]）；产量不带
    assert ws.cell(row=7, column=headers.index("Peak Year Spend") + 1).number_format == "#,##0\\ [$€-407]"  # 案例级：合并区首行
    assert ws.cell(row=8, column=headers.index("Average Volume") + 1).number_format == "#,##0"
    # 件价最多 4 位小数；模具费是整笔金额，显示 2 位
    assert ws.cell(row=8, column=headers.index("CQA PC Price") + 1).number_format == "#,##0.00##\\ [$€-407]"
    for name in ("CQA Tooling", "Rec. Supplier Tooling", "Tooling Saving vs CQA"):
        assert ws.cell(row=8, column=headers.index(name) + 1).number_format == "#,##0.00\\ [$€-407]", name
    assert "ISO calendar weeks (KW)" in ws["A2"].value and "7-day" not in ws["A2"].value
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref.startswith("A4:")


def test_project_type_and_supplier_are_written_per_part_number_row():
    rows = [dict(ROW_A, project="MBEAL", sourcingType="New", recommendedSupplier="Supplier A"),
            dict(ROW_B, project="ACR8", sourcingType="GCS", recommendedSupplier="Supplier B")]
    case = _multi(partNumbers=rows, project="MBEAL, ACR8", sourcingType="New, GCS", recommendedSupplier="Supplier A, Supplier B")
    ws = _sheet(exports.build("agenda", [case], "2026-KW39", WHEN), "Agenda")
    headers, out = _rows(ws)
    assert [(r["Project / Product"], r["Type"], r["Recommended Supplier(s)"]) for r in out] == [
        ("MBEAL", "New", "Supplier A"), ("ACR8", "GCS", "Supplier B")]
    merged = {str(m) for m in ws.merged_cells.ranges}
    for name in ("Type", "Project / Product", "Recommended Supplier(s)"):
        L = exports.get_column_letter(headers.index(name) + 1)
        assert f"{L}5:{L}6" not in merged, name
    for name in ("Peak Year Spend", "LT Spend"):
        L = exports.get_column_letter(headers.index(name) + 1)
        assert f"{L}5:{L}6" in merged, name


def test_savings_are_excel_formulas_following_the_template():
    ws = _sheet(exports.build("agenda", [_case()], "2026-KW39", WHEN), "Agenda")
    headers, rows = _rows(ws)
    L = {h: exports.get_column_letter(i) for i, h in enumerate(headers, start=1)}
    assert rows[0]["Average Year Saving vs CQA"] == f"=({L['Rec. Supplier Landed']}5-{L['CQA PC Price']}5)*{L['Average Volume']}5"
    assert rows[0]["LT Saving vs CQA"] == f"=({L['Rec. Supplier Landed']}5-{L['CQA PC Price']}5)*{L['LT Volume']}5"
    assert rows[0]["Tooling Saving vs CQA"] == f"={L['Rec. Supplier Tooling']}5-{L['CQA Tooling']}5"
    assert ws.cell(row=5, column=headers.index("Average Year Saving vs CQA") + 1).data_type == "f"


def test_formula_like_text_is_stored_as_text_not_formula():
    ws = _sheet(exports.build("agenda", [_case()], "2026-KW39", WHEN), "Agenda")
    headers, _ = _rows(ws)
    cell = ws.cell(row=5, column=headers.index("Comment") + 1)
    assert cell.value == "=SUM(A1)" and cell.data_type == "s"


def test_agenda_for_empty_week_still_has_headers():
    ws = _sheet(exports.build("agenda", [], "2026-KW40", WHEN), "Agenda")
    headers, rows = _rows(ws)
    assert headers[:2] == ["#", "KW"] and rows == []


def test_minutes_has_decision_notes_and_one_action_row_per_task():
    data = exports.build("minutes", [_case(), _case(caseNumber=3, swatId="SWAT-3", meetingDecision="PENDING", actionStatus=None, followUps=[],
                                                    committeeDiscussion="Pending for Sourcing Committee Review")], "2026-KW39", WHEN)
    minutes = _sheet(data, "Minutes")
    assert minutes["A1"].value == "Sourcing Committee Meeting Minutes — 2026-KW39" and minutes["A1"].font.color.rgb.endswith("375623")
    assert minutes.sheet_properties.tabColor.rgb.endswith("375623")
    headers, rows = _rows(minutes)
    assert headers.index("Sourcing Decision") == headers.index("Comment") + 1
    assert rows[0]["Sourcing Decision"] == "Approved with conditions" and rows[0]["Action Status"] == "Open"  # 决议按模板文案
    assert rows[0]["Notes / Task / Resp. / Due Date"] == ("Approved subject to consignment.\n"
                                                           "1. Negotiate consignment — SP ABC, due 2027-02-18 [Open]\n"
                                                           "2. Send BP — SP ABC, due 2026-10-01 [Closed]")
    assert rows[1]["Sourcing Decision"] == "Pending" and rows[1]["Action Status"] == "" and rows[1]["Notes / Task / Resp. / Due Date"] == ""
    actions = _sheet(data, "Actions")
    _, a_rows = _rows(actions)
    assert [r["Task"] for r in a_rows] == ["Negotiate consignment", "Send BP"]
    assert a_rows[0]["Category"] == "Supplier strategy, Saving" and a_rows[0]["Reminder to"] == "sp@example.com"
    assert a_rows[1]["Status"] == "Closed" and a_rows[1]["Reminder to"] == ""
    assert a_rows[0]["#"] == 1 and a_rows[1]["#"] == 2
    assert a_rows[0]["Sourcing Decision"] == "Approved with conditions"
    # 到期日是真日期单元格（Excel 可按日期筛选排序）
    due = actions.cell(row=5, column=[c.value for c in actions[4]].index("Due Date") + 1)
    assert due.value.date() == date(2027, 2, 18) and due.number_format == "yyyy-mm-dd"


def test_meeting_date_is_a_real_date_when_known_and_legacy_usd_amounts_carry_the_dollar_sign():
    ws = _sheet(exports.build("agenda", [_case(meetingDateISO="2026-09-30"), _case(caseNumber=3, swatId="SWAT-3", meetingDateISO="")],
                              "2026-KW39", WHEN), "Agenda")
    headers = [c.value for c in ws[4]]
    col = headers.index("Meeting Date") + 1
    assert ws.cell(row=5, column=col).value.date() == date(2026, 9, 30) and ws.cell(row=5, column=col).number_format == "yyyy-mm-dd"
    assert ws.cell(row=6, column=col).value == "Wed, Sep 30, 2026"  # 没有 ISO 日期的旧记录回退到原文本
    old = _sheet(exports.build("agenda", [logic.mk_case({"swatId": "SWAT-OLD", "peakYearSpend": 1500000}, [])], "2026-KW01", WHEN), "Agenda")
    assert old.cell(row=5, column=headers.index("Peak Year Spend") + 1).number_format == "#,##0\\ [$$-409]"


def test_late_registration_is_marked():
    ws = _sheet(exports.build("agenda", [_case(lateRegistration={"reason": "Customer escalation", "deadline": "x"})], "2026-KW39", WHEN), "Agenda")
    _, rows = _rows(ws)
    assert rows[0]["Late Registration"] == "Yes — Customer escalation"


def test_legacy_usd_case_without_commercial_fields_exports_blank_prices():
    ws = _sheet(exports.build("agenda", [logic.mk_case({"swatId": "SWAT-OLD", "peakYearSpend": 1500000}, [])], "2026-KW01", WHEN), "Agenda")
    _, rows = _rows(ws)
    assert rows[0]["Currency"] == "USD" and rows[0]["CQA PC Price"] == "" and rows[0]["Tooling Payment"] == ""


@pytest.mark.parametrize("kind,sheet,name", [("agenda", "Agenda", "Sourcing_Committee_Agenda_{label}.xlsx"),
                                             ("minutes", "Minutes", "Sourcing_Committee_Minutes_{label}.xlsx")])
def test_export_endpoint_streams_xlsx_for_the_week(api_client, kind, sheet, name):
    all_cases = api_client.get("/api/bootstrap").json()["cases"]
    week = max(c["weekNum"] for c in all_cases)  # 演示数据里最后一周
    r = api_client.get(f"/api/exports/{kind}", params={"week": week})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert f'filename="{name.format(label=logic.wk(week))}"' in r.headers["content-disposition"]
    ws = _sheet(r.content, sheet)
    _, rows = _rows(ws)
    seeded = [c for c in all_cases if c["weekNum"] == week]
    assert len(rows) == sum(len(logic.part_number_lines(c)) for c in seeded) > 0
    assert [r["#"] for r in rows if r["#"] != ""] == sorted(c["caseNumber"] for c in seeded)


def test_database_export_takes_the_selected_case_ids(api_client):
    cases = api_client.get("/api/bootstrap").json()["cases"]
    picked = cases[:3]
    r = api_client.post("/api/exports/minutes", json={"ids": [c["id"] for c in picked] + ["C9999"]})
    assert r.status_code == 200
    assert 'filename="Sourcing_Committee_Minutes_Database_2026-09-24.xlsx"' in r.headers["content-disposition"]
    ws = _sheet(r.content, "Minutes")
    assert ws["A1"].value == "Sourcing Committee Meeting Minutes — Database export, 3 cases (Sep 24, 9:00 AM)"
    _, rows = _rows(ws)
    assert [r["SWAT ID / Bundle ID"] for r in rows if r["SWAT ID / Bundle ID"]] == [c["swatId"] for c in sorted(picked, key=lambda c: (c["weekNum"], c["caseNumber"], c["swatId"]))]
    assert api_client.post("/api/exports/agenda", json={"ids": []}).status_code == 200  # 空选择：只有表头
    assert api_client.post("/api/exports/agenda", json={"ids": ["x"] * 5001}).status_code == 422


def test_export_endpoint_rejects_unknown_kind_and_bad_week(api_client):
    assert api_client.get("/api/exports/pdf", params={"week": 39}).status_code == 404
    assert api_client.post("/api/exports/pdf", json={"ids": []}).status_code == 404
    assert api_client.get("/api/exports/agenda", params={"week": 0}).status_code == 422
    assert api_client.get("/api/exports/agenda").status_code == 422
