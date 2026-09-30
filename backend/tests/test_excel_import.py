"""v3 Phase-17：周会 Excel（委员会模板的 Agenda / MM）导入成案例。

用 openpyxl 现造一份和 SharePoint 上真实文件同样式的工作簿：第 1 行表头、案例级单元格纵向合并、
零件按行、金额是 "439K Euro" 这种文本、决议是自由文案。真实文件的 64 份样本另外在本机上跑过一遍。
"""

import io
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from app import excel_import as ei
from app import logic

HEADERS = ["#", "Region", "Type", "Project/ Product", "Commodity", "Parent PF Code", "SWAT ID / Bundle ID",
           "Part Description", "PN", "Recommended Supplier (s)", "Presenter", "Decision Level", "Peak Year Spend Euro",
           "LT Spend\nEuro", "Comment", "Sourcing Decision", "Notes/Task/ Resp./ Due Date", "CQA PC PRICE",
           "Rec. Supplier Landed", "Average Volume", "LT volume", "Average Year Saving vs CQA", "LT saving",
           "CQA Tooling", "Rec. Supplier Tooling", "Tooling Saving vs CQA"]
COL = {h.split("\n")[0]: i + 1 for i, h in enumerate(HEADERS)}
FX = {"basis": "OP", "perEur": {"USD": "1.17", "CNY": "8.3"}}


def _workbook(rows: list[dict], merges: list[str], sheet="KW39, 2026") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(HEADERS)
    for r, row in enumerate(rows, start=2):
        for header, value in row.items():
            ws.cell(row=r, column=COL[header], value=value)
    for rng in merges:
        ws.merge_cells(rng)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 一份贴近真实文件的样本：案例 1 是三个零件的 bundle（金额按行），案例 2 单零件（外币、offline 决议），
# 案例 3 的 '#' 合并到第二行，第二行是同一零件的 NA 变体（自己的价格）
SAMPLE_ROWS = [
    {"#": 1, "Region": "EU, NA", "Type": "New", "Project/ Product": "MBEAL", "Commodity": "CH01", "Parent PF Code": "CH01016\nMedium Precision",
     "SWAT ID / Bundle ID": "B/2025/00361\n110000034017", "Part Description": "ECU Housing, WMS", "PN": "R003B136A",
     "Recommended Supplier (s)": "Ningbo Deke", "Presenter": "Sandra Novegil", "Decision Level": "L2",
     "Peak Year Spend Euro": "439K Euro", "LT Spend": "2,5Mio Euro", "Comment": "1. CN volume not yet nominated",
     "Sourcing Decision": "Approved with conditions", "Notes/Task/ Resp./ Due Date": "1. To check cleanliness --BUP Rafal",
     "CQA PC PRICE": "0,64 Euro/pc", "Rec. Supplier Landed": "0,74 Euro/pc", "Average Volume": 378979, "LT volume": 3410814,
     "CQA Tooling": 277000, "Rec. Supplier Tooling": 4352},
    {"Type": "New", "Project/ Product": "MBEAL", "SWAT ID / Bundle ID": "B/2025/00361\n110000034018", "Part Description": "ECU Housing, w/o WMS",
     "PN": "R003H147A", "Peak Year Spend Euro": "139K Euro", "LT Spend": "873K Euro", "CQA PC PRICE": "0,64 Euro/pc",
     "Rec. Supplier Landed": "0,66 Euro/pc", "Average Volume": 165280, "LT volume": 1322241, "CQA Tooling": 0, "Rec. Supplier Tooling": "138,908 Euro"},
    {"Type": "GCS", "Project/ Product": "ACR8", "SWAT ID / Bundle ID": "B/2025/00361\n110000034019", "Part Description": "ECU Cover",
     "PN": "R004D369A", "Peak Year Spend Euro": "139K Euro", "LT Spend": "852K Euro", "CQA PC PRICE": "0,12 Euro/pc",
     "Rec. Supplier Landed": "0,18 Euro/pc", "Average Volume": 525895, "LT volume": 4733055},
    {"#": 2, "Region": "AP", "Type": "Carry Over", "Project/ Product": "VW CMP21", "Commodity": "EA05", "Parent PF Code": "EA05001",
     "SWAT ID / Bundle ID": "110000035117", "Part Description": "Locking Disc Alu", "PN": "R003H729B",
     "Recommended Supplier (s)": "Shenzhen XLX", "Presenter": "Jordan ZHAO", "Decision Level": "L4",
     "Peak Year Spend Euro": "11.7K USD", "LT Spend": "61K Euro Euro", "Sourcing Decision": "Approved offline KW38",
     "CQA PC PRICE": "12.5 CNY", "Rec. Supplier Landed": "0.2108 Euro/pc\n1,75 RMB/pc", "Average Volume": "28,163", "LT volume": 253464,
     "CQA Tooling": "39,157 Euro\n325.000 RMB", "Rec. Supplier Tooling": "N/A"},
    {"#": 3, "Region": "EU", "Type": "New Part", "Project/ Product": "SPR6 Programs", "Commodity": "FF04", "Parent PF Code": "FF0401101",
     "SWAT ID / Bundle ID": "B/2025/00228", "Part Description": "SPR6.1 Locking Pawl", "PN": "R000R613A",
     "Recommended Supplier (s)": "Carbosint", "Presenter": "Daria W.", "Decision Level": "Level 2",
     "Peak Year Spend Euro": "1.1Mio Euro", "LT Spend": "7,5 Mio Euro", "Sourcing Decision": "Not approved, \nPending actions",
     "Notes/Task/ Resp./ Due Date": "1. Request ASIMCO to quote", "CQA PC PRICE": 0.12, "Rec. Supplier Landed": 0.2087,
     "Average Volume": 4610942, "LT volume": 32276597, "CQA Tooling": 40000, "Rec. Supplier Tooling": 284000},
    {"Region": "NA", "CQA PC PRICE": 0.23433, "Rec. Supplier Landed": 0.21166, "Average Volume": 1068308, "LT volume": 9614770},
    {"#": 4, "Region": "AP", "Type": "New", "Project/ Product": "Agenda only", "SWAT ID / Bundle ID": "TBD in supplyon",
     "Part Description": "Bracket", "PN": "R004C217A", "Recommended Supplier (s)": "AMP", "Presenter": "L. Hu", "Decision Level": "L3",
     "Peak Year Spend Euro": "4 < 750k", "LT Spend": 149000},
]
# 案例 1 的案例级列合并 A2:A4 …；案例 3 合并 A6:A7 与零件列 H6:H7 / I6:I7（NA 行是同一零件）
SAMPLE_MERGES = ["A2:A4", "B2:B4", "G2:G2", "J2:J4", "K2:K4", "L2:L4", "O2:O4", "P2:P4", "Q2:Q4",
                 "A6:A7", "C6:C7", "D6:D7", "G6:G7", "H6:H7", "I6:I7", "J6:J7", "K6:K7", "L6:L7", "M6:M7", "N6:N7", "P6:P7", "Q6:Q7"]


@pytest.fixture
def sample() -> bytes:
    return _workbook(SAMPLE_ROWS, SAMPLE_MERGES)


# ---------- 文件名 ----------

@pytest.mark.parametrize("name,kw,day", [
    ("KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", 39, date(2026, 9, 23)),
    ("KW37, 10.09.2025_SBS Sourcing Alignment Committee_MM updated.xlsx", 37, date(2025, 9, 10)),
    ("20260128 sourcing alignment meeting_KW05_MM.xlsx", 5, date(2026, 1, 28)),
    ("RB WK15. Apr 08 2026. sourcing alignment committee MM.xlsx", 15, date(2026, 4, 8)),
    ("KW39 24 26.09.2025_SBS Sourcing Alignment Committee_MM.xlsx", 39, date(2025, 9, 24)),  # 26.09 是准备日：以 KW 为准取周三
    ("KW36 31.08.2026_SBS Sourcing Alignment Committee_MM.xlsx", 36, date(2026, 9, 2)),
    ("KW46_SBS Sourcing Alignment Committee_MM.xlsx", 46, None),  # 没有年份：日期由导入表单填
    ("KW46, 12.11.2025", 46, date(2025, 11, 12)),  # SharePoint 周文件夹名
    ("random.xlsx", None, None),
])
def test_meeting_date_from_file_name(name, kw, day):
    assert ei.meeting_from_name(name) == (kw, day)


# ---------- 金额与文案 ----------

@pytest.mark.parametrize("text,expected", [
    ("439K Euro", 439_000), ("2,5Mio Euro", 2_500_000), ("1.03 Mio Euro", 1_030_000), ("198.7K EUR", 198_700),
    ("1.3M EUR", 1_300_000), ("0.132Mio Euro", 132_000), ("61K Euro Euro", 61_000), ("228K Euro\n251K USD", 228_000),
    ("1,03 Mio Euro\n8,15 Mio RMB", 1_030_000), ("4,66 – 7,43 M€", 7_430_000), ("60,000 Euro", 60_000),
    ("7.500 Euro", 7_500), ("253,164.56 EUR", "253164.56"), ("5k", 5_000), ("2.27 Mio", 2_270_000),
    (12466, 12466), (222089.79, "222089.79"), ("0 Euro", 0),
])
def test_spend_text_is_read_in_euro(text, expected):
    amount, note = ei.parse_amount(text, FX)
    assert amount == ei.Decimal(str(expected)) and note is None


def test_foreign_amounts_convert_with_dashboard_rates_and_warn_without():
    amount, note = ei.parse_amount("251K USD", FX)
    assert amount == ei.Decimal("214529.9145") and "USD converted" in note
    amount, note = ei.parse_amount("12.5 CNY", FX, price=True)
    assert amount == ei.Decimal("1.5060") and "CNY converted" in note
    amount, note = ei.parse_amount("2,000,000 RMB", {"perEur": {}})
    assert amount is None and "no CNY rate" in note
    for junk in ("4 < 750k", "3 < 1.5mio", "see PPT"):
        amount, note = ei.parse_amount(junk, FX)
        assert amount is None and "cannot read" in note
    assert ei.parse_amount("N/A", FX) == (None, None) and ei.parse_amount("", FX) == (None, None)


@pytest.mark.parametrize("text,expected", [
    ("0,64 Euro/pc", "0.64"), ("0,3258 Euro/PC    0,3584 USD/PC", "0.3258"), ("29,601 RMB/pc\n2,746 Euro/pc", "2.746"),
    ("0.81575 Euro/pc\n0,89733 USD/pc", "0.81575"), ("581,15 Euro/kg", "581.15"), ("1.6", "1.6"),
])
def test_piece_prices_treat_commas_as_decimal_points(text, expected):
    amount, _ = ei.parse_amount(text, FX, price=True)
    assert amount == ei.Decimal(expected)


@pytest.mark.parametrize("text,expected", [
    ("Approved", "APPROVED"), ("Approved offline KW38", "APPROVED"), ("Offline Approved KW17", "APPROVED"),
    ("Approved with conditions", "CONDITIONAL APPROVAL"), ("Conditionally Approved with actions", "CONDITIONAL APPROVAL"),
    ("Approved with action", "CONDITIONAL APPROVAL"), ("Rejected.", "REJECTED"), ("Not approved, \nPending actions", "REJECTED"),
    ("Offline reviewed KW38, Rejected, To be followed up", "REJECTED"), ("To be followed up", "PENDING"),
    ("Pending Offline Approval", "PENDING"), ("On hold for further clarification", "PENDING"), ("n/a", "PENDING"), ("", "PENDING"),
])
def test_decision_text_is_classified(text, expected):
    assert ei.classify_decision(text) == expected


def test_type_level_region_and_ids_are_normalised():
    assert [ei.normalize_type(t) for t in ("New", "NEW", "New Part", "Family part", "Carry Over", "CO", "C/O", "Carry-over", "GCS", "CO GCS", "Make vs buy")] == \
        ["New", "New", "New", "New", "C/O", "C/O", "C/O", "C/O", "GCS", "GCS", "Make vs buy"]
    assert [ei.normalize_level(t) for t in ("L2", "L3", "Level 4", "2", "")] == ["Level 2", "Level 3", "Level 4", "Level 2", ""]
    assert ei.normalize_region("EU, NA") == "EU + NA" and ei.normalize_region("EU (AP)") == "AP + EU" and ei.normalize_region("EU ") == "EU"
    assert ei.merge_regions(["EU", "NA", "EU"]) == "EU + NA"
    assert ei.split_ids("B/2025/00361\n110000034017") == ("B/2025/00361", "110000034017")
    assert ei.split_ids("110000034017 / B/2025/00361") == ("B/2025/00361", "110000034017")
    assert ei.split_ids("110000033111") == ("110000033111", "110000033111")
    assert ei.split_ids("TBD in supplyon") == ("TBD in supplyon", "")
    assert ei.split_part_numbers("R002P856C, R002P857C, R002P858C") == ["R002P856C", "R002P857C", "R002P858C"]
    assert ei.split_part_numbers("R002G246A\n") == ["R002G246A"]
    assert ei.split_part_numbers("Wiring assy R002B221A / bracket") == ["Wiring assy R002B221A / bracket"]


# ---------- 整份工作簿 ----------

def test_sample_workbook_becomes_cases_with_bundle_rows(sample):
    report = ei.parse_workbook(sample, "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", FX, imported_by="admin@zf.com", when="now")
    assert (report["weekNum"], report["meetingYear"], report["meetingDateISO"], report["meetingDateLabel"]) == (39, 2026, "2026-09-23", "Wed, Sep 23, 2026")
    cases = report["cases"]
    assert [c["caseNumber"] for c in cases] == [1, 2, 3, 4] and report["rows"] == 7
    one, two, three, four = cases
    # 案例 1：bundle 号做案例号，三行零件各带自己的项目 / 类型 / 供应商（留空的沿用案例级），金额逐行相加
    assert one["swatId"] == "B/2025/00361" and one["region"] == "EU + NA" and one["decisionLevel"] == "Level 2"
    assert [(p["partNumber"], p["project"], p["sourcingType"], p["recommendedSupplier"], p.get("sourcingId")) for p in one["partNumbers"]] == [
        ("R003B136A", "MBEAL", "New", "Ningbo Deke", "110000034017"), ("R003H147A", "MBEAL", "New", "Ningbo Deke", "110000034018"),
        ("R004D369A", "ACR8", "GCS", "Ningbo Deke", "110000034019")]
    assert (one["project"], one["sourcingType"]) == ("MBEAL, ACR8", "New, GCS")
    assert (one["peakYearSpend"], one["lifetimeSpend"], one["spendCurrency"]) == (717_000, 4_225_000, "EUR")
    assert one["partNumbers"][0]["pcPriceCQA"] == "0.64" and one["partNumbers"][1]["supplierToolingCost"] == "138908"
    assert one["partNumbers"][0]["averageVolume"] == 378979 and one["pcPriceCQA"] is None  # 多零件：顶层不放件价
    assert one["toolingPayment"] == ""  # 商务区块按"字段存在"显示，导入的案例要带上（空）
    assert (one["meetingDecision"], one["decisionText"], one["caseStatus"]) == ("CONDITIONAL APPROVAL", "Approved with conditions", "Approved")
    assert one["committeeDiscussion"] == "1. To check cleanliness --BUP Rafal" and one["submitterComments"] == "1. CN volume not yet nominated"
    assert one["cluster"] == "CH01" and one["parentPF"] == "CH01016 Medium Precision" and one["presenter"] == "Sandra Novegil"
    assert one["source"] == {"kind": "excel", "file": "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", "row": 2,
                             "importedAt": "now", "importedBy": "admin@zf.com"}
    assert one["meetingDateISO"] == "2026-09-23" and one["meetingYear"] == 2026 and one["weekNum"] == 39
    # 案例 2：单零件放顶层；外币按 Dashboard 汇率折算；"39,157 Euro" 是三万九；N/A 为空
    assert two["partNumbers"] is None and two["partNumber"] == "R003H729B" and two["sourcingType"] == "C/O"
    assert (two["peakYearSpend"], two["lifetimeSpend"]) == (10_000, 61_000)  # 11.7K USD / 1.17
    assert (two["pcPriceCQA"], two["supplierPriceLanded"], two["toolingCQA"], two["supplierToolingCost"]) == ("1.506", "0.2108", "39157", None)
    assert two["averageVolume"] == 28163 and two["meetingDecision"] == "APPROVED" and two["decisionText"] == "Approved offline KW38"
    # 案例 3：'#' 合并两行；第二行是同一零件的 NA 变体 → 两行同零件号、带各自区域与价格；案例区域取并集
    assert three["region"] == "EU + NA" and three["decisionLevel"] == "Level 2" and three["meetingDecision"] == "REJECTED"
    assert [p["partNumber"] for p in three["partNumbers"]] == ["R000R613A", "R000R613A"]
    assert three["partNumbers"][1]["region"] == "NA" and three["partNumbers"][1]["pcPriceCQA"] == "0.2343"  # 件价统一存 4 位小数
    assert (three["peakYearSpend"], three["lifetimeSpend"]) == (1_100_000, 7_500_000)  # 合并单元格只算一次
    assert three["caseStatus"] == "Resubmission Required"
    # 案例 4：议程行没有决议 → Pending；读不懂的金额记 0 并给警告；SWAT 是自由文本原样保留
    assert four["meetingDecision"] == "PENDING" and four["committeeDiscussion"] == "Pending for Sourcing Committee Review"
    assert (four["peakYearSpend"], four["lifetimeSpend"]) == (0, 149_000) and four["swatId"] == "TBD in supplyon"
    assert any("Row 8 peak spend: cannot read amount" in w for w in report["warnings"])
    assert any("CNY converted" in w for w in report["warnings"])
    # 与 logic 的案例形状兼容：编辑用的零件行能拼出来
    assert logic.part_number_lines(two)[0]["recommendedSupplier"] == "Shenzhen XLX"


def test_workbook_needs_a_recognisable_header_and_a_meeting_date(sample):
    with pytest.raises(ei.TemplateError, match="meeting date"):
        ei.parse_workbook(sample, "KW46_SBS Sourcing Alignment Committee_MM.xlsx", FX)
    report = ei.parse_workbook(sample, "KW46_SBS Sourcing Alignment Committee_MM.xlsx", FX, meeting_date=date(2025, 11, 12))
    assert (report["weekNum"], report["meetingYear"]) == (46, 2025)
    wb = Workbook()
    wb.active.append(["Name", "Value"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ei.TemplateError, match="committee template"):
        ei.parse_workbook(buf.getvalue(), "KW39 23.09.2026_x.xlsx", FX)


def test_header_variants_and_extra_columns_are_tolerated():
    headers = ["Presenting Time", "#", "Region", "Type", "Project/ Product", "Commodity", "SWAT ID / Bundle ID", "Part Description", "PN",
               "Recommended Supplier (s)", "Presenter", "Decision Level", "Peak Year Spend", "LT Spend", "Sourcing Decision",
               "Notes/Task/ Resp./ Due Date"]  # KW37 2026 多了 Presenting Time；早期文件没有 Comment 与商务列
    wb = Workbook()
    ws = wb.active
    ws.append(["Some title row"])
    ws.append(headers)
    ws.append(["09:00", 1, "EU", "GCS", "P1", "SC", "110000033111", "Motor", "R003C494A", "Wuxi MI", "C. Wang", "L3", "173K Euro", "1.03 Mio Euro", "Approved", ""])
    buf = io.BytesIO()
    wb.save(buf)
    report = ei.parse_workbook(buf.getvalue(), "KW37 09.09.2026_SBS Sourcing Alignment Committee_Updated Agenda 3.xlsx", FX)
    case = report["cases"][0]
    assert (case["swatId"], case["sourcingType"], case["lifetimeSpend"], case["meetingDecision"]) == ("110000033111", "GCS", 1_030_000, "APPROVED")
    assert case["pcPriceCQA"] is None and case["submitterComments"] == ""


# ---------- 接口：管理员专属，先预览后写库，同一周重导替换 ----------

@pytest.fixture
def admin(api_client):
    r = api_client.post("/api/auth/login", json={"email": "admin@zf.com", "password": "test-admin"})
    return {"X-Session-Token": r.json()["token"]}


def _post(client, data: bytes, name: str, headers: dict, **params):
    return client.post("/api/import/excel", params={"name": name, **params}, content=data, headers={
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", **headers})


def test_import_endpoint_previews_then_commits_and_replaces_the_same_week(api_client, admin, sample):
    name = "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx"
    assert _post(api_client, sample, name, {"X-Session-Token": ""}).status_code == 401
    assert _post(api_client, sample, name, {}).status_code == 403  # 普通用户
    before = len(api_client.get("/api/bootstrap").json()["cases"])
    # 预览：不写库
    r = _post(api_client, sample, name, admin)
    assert r.status_code == 200, r.text
    p = r.json()
    assert (p["cases"], p["rows"], p["weekNum"], p["meetingYear"], p["replaces"], p["committed"]) == (4, 7, 39, 2026, 0, False)
    assert p["preview"][0] == {"caseNumber": 1, "swatId": "B/2025/00361", "partDescription": "ECU Housing, WMS", "parts": 3,
                               "region": "EU + NA", "decision": "CONDITIONAL APPROVAL", "lifetimeSpend": 4_225_000}
    assert "snapshot" not in p and len(api_client.get("/api/bootstrap").json()["cases"]) == before
    # 写库
    r = _post(api_client, sample, name, admin, commit="true", sourceUrl="https://sp.example.com/KW39")
    assert r.status_code == 200, r.text
    assert r.json()["committed"] is True and r.json()["added"] == 4
    cases = [c for c in r.json()["snapshot"]["cases"] if c.get("source", {}).get("kind") == "excel"]
    assert len(cases) == 4 and len(r.json()["snapshot"]["cases"]) == before + 4
    decided = next(c for c in cases if c["caseNumber"] == 1)
    assert decided["sourcingPresentationLink"] == "https://sp.example.com/KW39" and decided["source"]["importedBy"] == "admin@zf.com"
    assert next(c for c in cases if c["caseNumber"] == 4)["sourcingPresentationLink"] == ""  # Pending 不挂链接
    # 同一周再导（换了文件名也一样）：替换上次导入的，其它案例不动
    r = _post(api_client, _workbook(SAMPLE_ROWS[3:4], []), "KW39 23.09.2026_SBS Sourcing Alignment Committee_Agenda.xlsx", admin)
    assert r.json()["replaces"] == 4
    r = _post(api_client, _workbook(SAMPLE_ROWS[3:4], []), "KW39 23.09.2026_SBS Sourcing Alignment Committee_Agenda.xlsx", admin, commit="true")
    assert r.status_code == 200 and r.json()["added"] == 1
    snap = api_client.get("/api/bootstrap").json()
    imported = [c for c in snap["cases"] if c.get("source", {}).get("kind") == "excel"]
    assert len(imported) == 1 and len(snap["cases"]) == before + 1
    assert len({c["id"] for c in snap["cases"]}) == len(snap["cases"])
    # 另一周不受影响；会议日期可以手填
    r = _post(api_client, sample, "KW46_SBS Sourcing Alignment Committee_MM.xlsx", admin, commit="true", meetingDate="2025-11-12")
    assert r.status_code == 200 and r.json()["weekNum"] == 46 and r.json()["meetingYear"] == 2025
    assert len([c for c in api_client.get("/api/bootstrap").json()["cases"] if c.get("source", {}).get("kind") == "excel"]) == 5


def test_import_endpoint_rejects_bad_input_readably(api_client, admin, sample):
    assert _post(api_client, sample, "KW46_SBS Sourcing Alignment Committee_MM.xlsx", admin).status_code == 422  # 猜不出日期
    assert _post(api_client, sample, "deck.pptx", admin).status_code == 422
    assert _post(api_client, b"not a workbook", "KW39 23.09.2026_x.xlsx", admin).status_code == 422
    assert _post(api_client, sample, "KW39 23.09.2026_x.xlsx", admin, meetingDate="2026-13-01").status_code == 422
    assert _post(api_client, sample, "KW39 23.09.2026_x.xlsx", admin, sourceUrl="javascript:alert(1)").status_code == 422
    empty = _workbook([], [])
    r = _post(api_client, empty, "KW39 23.09.2026_x.xlsx", admin, commit="true")
    assert r.status_code == 422 and "No cases" in r.json()["detail"]


def test_imported_case_can_be_edited_without_a_presentation_link(api_client, admin, sample):
    name = "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx"
    r = _post(api_client, sample, name, admin, commit="true")
    case = next(c for c in r.json()["snapshot"]["cases"] if c.get("source", {}).get("kind") == "excel" and c["caseNumber"] == 2)
    body = {"partNumbers": [{"partNumber": case["partNumber"], "partDescription": case["partDescription"], "project": "VW CMP21",
                             "sourcingType": "C/O", "recommendedSupplier": "Shenzhen XLX"}],
            "region": case["region"], "meetingDecision": "APPROVED", "followUps": [], "committeeDiscussion": "Confirmed in the Cockpit."}
    r = api_client.put(f"/api/cases/{case['id']}", json=body, headers=admin)
    assert r.status_code == 200, r.text  # 系统内登记的案例这里会因为缺演示文稿链接被拒
    assert r.json()["case"]["committeeDiscussion"] == "Confirmed in the Cockpit." and r.json()["case"]["source"]["kind"] == "excel"


# ---------- 随程序发布的历史文件库：下拉框选文件导入 ----------

@pytest.fixture
def library(tmp_path, monkeypatch, sample):
    root = tmp_path / "history"
    (root / "2025" / "KW46, 12.11.2025").mkdir(parents=True)
    (root / "2026" / "KW39, 23.09.2026").mkdir(parents=True)
    (root / "2025" / "KW46, 12.11.2025" / "KW46_SBS Sourcing Alignment Committee_MM.xlsx").write_bytes(sample)  # 日期在文件夹名
    (root / "2026" / "KW39, 23.09.2026" / "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx").write_bytes(sample)
    (root / "2026" / "KW39, 23.09.2026" / "~$KW39 lock.xlsx").write_bytes(b"lock")  # Excel 的临时锁文件跳过
    monkeypatch.setattr(ei, "LIBRARY_DIR", root)
    return root


def test_library_lists_stored_files_with_dates_and_imports_by_path(api_client, admin, library):
    assert api_client.get("/api/import/library", headers={"X-Session-Token": ""}).status_code == 401
    assert api_client.get("/api/import/library").status_code == 403
    r = api_client.get("/api/import/library", headers=admin)
    assert r.status_code == 200
    listed = r.json()["files"]
    assert [(f["meetingYear"], f["weekNum"], f["meetingDateISO"], f["imported"]) for f in listed] == [
        (2025, 46, "2025-11-12", False), (2026, 39, "2026-09-23", False)]
    assert listed[0]["path"] == "2025/KW46, 12.11.2025/KW46_SBS Sourcing Alignment Committee_MM.xlsx"
    # 预览不写库；提交后列表标记 imported
    r = api_client.post("/api/import/library", json={"path": listed[0]["path"]}, headers=admin)
    assert r.status_code == 200 and r.json()["cases"] == 4 and r.json()["weekNum"] == 46 and r.json()["committed"] is False
    before = len(api_client.get("/api/bootstrap").json()["cases"])
    r = api_client.post("/api/import/library", json={"path": listed[0]["path"], "commit": True, "sourceUrl": "https://sp.example.com/KW46"}, headers=admin)
    assert r.status_code == 200 and r.json()["added"] == 4
    assert len(api_client.get("/api/bootstrap").json()["cases"]) == before + 4
    imported = [f["imported"] for f in api_client.get("/api/import/library", headers=admin).json()["files"]]
    assert imported == [True, False]
    case = next(c for c in api_client.get("/api/bootstrap").json()["cases"] if c.get("source", {}).get("kind") == "excel" and c["caseNumber"] == 1)
    assert case["sourcingPresentationLink"] == "https://sp.example.com/KW46" and case["meetingYear"] == 2025
    # 再导一次替换，不重复
    r = api_client.post("/api/import/library", json={"path": listed[0]["path"], "commit": True}, headers=admin)
    assert r.json()["replaces"] == 4 and len(api_client.get("/api/bootstrap").json()["cases"]) == before + 4


def test_library_rejects_paths_outside_the_folder(api_client, admin, library, tmp_path):
    (tmp_path / "outside.xlsx").write_bytes(b"x")
    for bad in ("../outside.xlsx", "/etc/passwd", "2026/KW39, 23.09.2026/missing.xlsx", "2026", ""):
        r = api_client.post("/api/import/library", json={"path": bad}, headers=admin)
        assert r.status_code in (404, 422), bad
    assert api_client.post("/api/import/library", json={"path": "2026/KW39, 23.09.2026/KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx", "sourceUrl": "javascript:x"}, headers=admin).status_code == 422


def test_shipped_library_is_complete():
    # 仓库里带的 36 份周文件（2025 KW36 – 2026 KW39）每份都能读出日期
    listed = ei.library_files()
    assert len(listed) == 36 and all(f["meetingDateISO"] for f in listed)
    assert (listed[0]["meetingYear"], listed[0]["weekNum"]) == (2025, 36) and (listed[-1]["meetingYear"], listed[-1]["weekNum"]) == (2026, 39)


# ---------- 命令行批量导入 ----------

async def test_cli_imports_a_folder_and_reads_the_date_from_the_folder_name(async_client, tmp_path, sample, capsys):
    from app import main, migrations
    week = tmp_path / "KW46, 12.11.2025"
    week.mkdir()
    (week / "KW46_SBS Sourcing Alignment Committee_MM.xlsx").write_bytes(sample)  # 文件名没有日期 → 用文件夹名
    (tmp_path / "KW39 23.09.2026_SBS Sourcing Alignment Committee_MM.xlsx").write_bytes(sample)
    (tmp_path / "junk.xlsx").write_bytes(b"nope")
    assert await migrations.run(["import-excel", str(tmp_path)], main.SessionLocal) == 1  # junk 失败 → 退出码 1，其余照常预览
    out = capsys.readouterr().out
    assert "DRY  KW46_SBS" in out and "2025-KW46" in out and "FAIL  junk.xlsx" in out
    before = len((await async_client.get("/api/bootstrap")).json()["cases"])
    (tmp_path / "junk.xlsx").unlink()
    assert await migrations.run(["import-excel", str(tmp_path), "--apply"], main.SessionLocal) == 0
    snap = (await async_client.get("/api/bootstrap")).json()
    imported = [c for c in snap["cases"] if c.get("source", {}).get("kind") == "excel"]
    assert len(imported) == 8 and len(snap["cases"]) == before + 8
    assert {(c["meetingYear"], c["weekNum"]) for c in imported} == {(2025, 46), (2026, 39)}
    # 再跑一次：同一周替换，总数不变
    assert await migrations.run(["import-excel", str(tmp_path), "--apply"], main.SessionLocal) == 0
    assert len((await async_client.get("/api/bootstrap")).json()["cases"]) == before + 8


def test_week_export_separates_years_once_history_is_imported(api_client, admin, sample):
    # 2025-KW39 与 2026-KW39 都导入后，按周导出默认只取当年（测试时钟 2026）的案例，year 参数可指定往年
    for name in ("KW39 23.09.2026_x_MM.xlsx", "KW39 24.09.2025_x_MM.xlsx"):
        assert _post(api_client, sample, name, admin, commit="true").status_code == 200

    def exported_rows(params):
        r = api_client.get("/api/exports/agenda", params=params, headers=admin)
        ws = load_workbook(io.BytesIO(r.content))["Agenda"]
        hdr = [c.value for c in ws[4]]
        return [row[hdr.index("Meeting Date")].value for row in ws.iter_rows(min_row=5) if row[hdr.index("PN")].value]
    this_year, last_year = exported_rows({"week": 39}), exported_rows({"week": 39, "year": 2025})
    assert len(this_year) == 7 and len(last_year) == 7 and this_year != last_year
    assert exported_rows({"week": 39, "year": 2024}) == []
